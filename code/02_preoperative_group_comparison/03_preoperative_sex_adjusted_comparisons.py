#!/usr/bin/env python3
"""Create article-ready sex-adjusted biomarker Table S6.

Input:
    data/preoperative_group_comparison.csv
Output:
    outputs/02_preoperative_group_comparison/table/TableS6_sex_adjusted_biomarkers.xlsx

The model formula, fixed endpoint scales, HC3 inference, Type III partial eta
squared, and one eight-endpoint BH-FDR family remain unchanged. Optional
residual diagnostics are displayed only and are never saved.
"""

from __future__ import annotations

import argparse
import math
import warnings
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

import numpy as np
import pandas as pd
from scipy import stats
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
import statsmodels.formula.api as smf
from statsmodels.regression.linear_model import RegressionResultsWrapper
from statsmodels.stats.anova import anova_lm
from statsmodels.stats.multitest import multipletests

PROJECT_DIR = Path(__file__).resolve().parents[2]
DATA_FILE = PROJECT_DIR / "data" / "preoperative_group_comparison.csv"
TABLE_DIR = PROJECT_DIR / "outputs" / "02_preoperative_group_comparison" / "table"


@dataclass(frozen=True)
class AncovaSpec:
    column: str
    label: str
    scale: Literal["original", "log"]


# Fixed after residual-distribution review. The optional diagnostics display
# candidate models but never change these assignments at runtime.
ANCOVA_SPECS: tuple[AncovaSpec, ...] = (
    AncovaSpec("iron_umol_l", "Fe (µmol/L)", "log"),
    AncovaSpec("hemoglobin_g_l", "Hb (g/L)", "original"),
    AncovaSpec("rbc_10e12_l", "RBC (×10¹²/L)", "log"),
    AncovaSpec("alt_u_l", "ALT (U/L)", "log"),
    AncovaSpec("ggt_u_l", "GGT (U/L)", "log"),
    AncovaSpec("hdl_c_mmol_l", "HDL-C (mmol/L)", "original"),
    AncovaSpec("creatinine_umol_l", "Creatinine (µmol/L)", "log"),
    AncovaSpec("uric_acid_umol_l", "Uric acid (µmol/L)", "log"),
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate Table S6 and optionally display ANCOVA residual diagnostics."
    )
    parser.add_argument(
        "--diagnostics",
        action="store_true",
        help=(
            "After writing Table S6, display residual diagnostics for both "
            "original- and log-scale candidate models. No diagnostic file is saved."
        ),
    )
    parser.add_argument(
        "--diagnostic-variables",
        nargs="+",
        metavar="COLUMN",
        help=(
            "Optional biomarker column names to display when --diagnostics is used. "
            "When omitted, all eight biomarkers are reviewed."
        ),
    )
    return parser.parse_args()



def normalise_group(series: pd.Series) -> pd.Series:
    values = series.astype("string").str.strip().str.lower()
    mapped = values.map({"non-be": 0, "non_be": 0, "nonbe": 0, "be": 1})
    unknown = sorted(values[mapped.isna() & values.notna()].dropna().unique().tolist())
    if unknown:
        raise ValueError(f"Unrecognised bes_group value(s): {unknown}")
    return mapped.astype("Int64")


def normalise_sex(series: pd.Series) -> pd.Series:
    values = series.astype("string").str.strip().str.lower()
    mapped = values.map({"female": 0, "f": 0, "0": 0, "male": 1, "m": 1, "1": 1})
    unknown = sorted(values[mapped.isna() & values.notna()].dropna().unique().tolist())
    if unknown:
        raise ValueError(f"Unrecognised sex value(s): {unknown}")
    return mapped.astype("Int64")


def load_data() -> pd.DataFrame:
    if not DATA_FILE.exists():
        raise FileNotFoundError(f"Input file not found: {DATA_FILE}")
    df = pd.read_csv(DATA_FILE)
    df.columns = [str(column).strip() for column in df.columns]
    df = df.replace(r"^\s*$", np.nan, regex=True)

    required = {"participant_id", "sex", "bes_group", *(spec.column for spec in ANCOVA_SPECS)}
    missing = sorted(required.difference(df.columns))
    if missing:
        raise ValueError(f"Missing required column(s): {missing}")

    df["participant_id"] = pd.to_numeric(df["participant_id"], errors="coerce")
    if df["participant_id"].isna().any() or df["participant_id"].duplicated().any():
        raise ValueError("participant_id must be complete, numeric, and unique.")

    df["group_be"] = normalise_group(df["bes_group"])
    df["sex_male"] = normalise_sex(df["sex"])
    for spec in ANCOVA_SPECS:
        df[spec.column] = pd.to_numeric(df[spec.column], errors="coerce")

    observed_groups = df["group_be"].value_counts().sort_index().to_dict()
    if len(df) != 122 or observed_groups != {0: 66, 1: 56}:
        raise AssertionError(
            f"Expected 122 participants (Non-BE=66, BE=56); observed {len(df)} and {observed_groups}."
        )
    if df["sex_male"].isna().any():
        raise ValueError("Sex must be complete for the sex-adjusted analysis.")

    return df.sort_values("participant_id", kind="mergesort").reset_index(drop=True)


def prepare_model_data(
    df: pd.DataFrame,
    spec: AncovaSpec,
    scale: Literal["original", "log"],
) -> pd.DataFrame:
    sub = df[[spec.column, "group_be", "sex_male"]].dropna().copy()
    if len(sub) < 4:
        raise ValueError(f"Insufficient complete observations for {spec.label}.")
    if sub["group_be"].nunique() != 2 or sub["sex_male"].nunique() != 2:
        raise ValueError(f"Both group and sex levels are required for {spec.label}.")

    if scale == "log":
        if bool((sub[spec.column] <= 0).any()):
            raise ValueError(f"Log-ANCOVA requires positive values for {spec.label}.")
        sub["model_outcome"] = np.log(sub[spec.column])
    else:
        sub["model_outcome"] = sub[spec.column]
    return sub


def fit_candidate_model(
    df: pd.DataFrame,
    spec: AncovaSpec,
    scale: Literal["original", "log"],
) -> tuple[pd.DataFrame, RegressionResultsWrapper]:
    sub = prepare_model_data(df, spec, scale)
    model = smf.ols("model_outcome ~ group_be + sex_male", data=sub).fit()
    return sub, model


def fit_endpoint(df: pd.DataFrame, spec: AncovaSpec) -> dict[str, object]:
    sub, model = fit_candidate_model(df, spec, spec.scale)

    robust = model.get_robustcov_results(cov_type="HC3")
    parameter_names = model.model.exog_names
    group_index = parameter_names.index("group_be")

    beta = float(robust.params[group_index])
    confidence_interval = np.asarray(robust.conf_int(alpha=0.05), dtype=float)[group_index]
    p_value = float(robust.pvalues[group_index])

    type3 = anova_lm(model, typ=3)
    group_ss = float(type3.loc["group_be", "sum_sq"])
    residual_ss = float(type3.loc["Residual", "sum_sq"])
    partial_eta_squared = group_ss / (group_ss + residual_ss)

    if spec.scale == "log":
        model_label = "log-ANCOVA-HC3"
        effect_label = "GMR"
        effect = math.exp(beta)
        ci_lower = math.exp(float(confidence_interval[0]))
        ci_upper = math.exp(float(confidence_interval[1]))
    else:
        model_label = "ANCOVA-HC3"
        effect_label = "AMD"
        effect = beta
        ci_lower = float(confidence_interval[0])
        ci_upper = float(confidence_interval[1])

    return {
        "indicator": spec.label,
        "column": spec.column,
        "n_complete": len(sub),
        "model": model_label,
        "model_scale": spec.scale,
        "effect_type": effect_label,
        "effect": float(effect),
        "ci_lower": float(ci_lower),
        "ci_upper": float(ci_upper),
        "p_value": p_value,
        "fdr_bh": np.nan,
        "partial_eta_squared": float(partial_eta_squared),
    }


def fmt_p(value: float) -> str:
    if not np.isfinite(value):
        return "—"
    return "<0.001" if value < 0.001 else f"{value:.3f}"


def fmt_eta(value: float) -> str:
    if not np.isfinite(value):
        return "—"
    return "<0.001" if value < 0.001 else f"{value:.3f}"


def fmt_effect(row: pd.Series) -> str:
    prefix = str(row["effect_type"])
    return (
        f"{prefix} {float(row['effect']):.3f} "
        f"({float(row['ci_lower']):.3f}, {float(row['ci_upper']):.3f})"
    )


def build_formatted_table(results: pd.DataFrame) -> pd.DataFrame:
    return pd.DataFrame(
        {
            "Indicator (unit)": results["indicator"],
            "N": results["n_complete"].astype(int),
            "Model": results["model"],
            "Adjusted effect (95% CI)": results.apply(fmt_effect, axis=1),
            "P value": results["p_value"].map(fmt_p),
            "q value (BH-FDR)": results["fdr_bh"].map(fmt_p),
            "Partial η²": results["partial_eta_squared"].map(fmt_eta),
        }
    )


ARTICLE_FONT = "Arial"
HEADER_SIDE = Side(style="thin", color="808080")
ROW_SIDE = Side(style="hair", color="D9D9D9")


def write_table_s6(table: pd.DataFrame) -> Path:
    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Table S6"
    worksheet.append(list(table.columns))
    for cell in worksheet[1]:
        cell.font = Font(name=ARTICLE_FONT, size=9, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=HEADER_SIDE)
    worksheet.row_dimensions[1].height = 38

    for row in table.itertuples(index=False, name=None):
        worksheet.append(list(row))
        row_index = worksheet.max_row
        for column_index, cell in enumerate(worksheet[row_index], start=1):
            cell.font = Font(name=ARTICLE_FONT, size=9)
            cell.alignment = Alignment(
                horizontal="left" if column_index == 1 else "center",
                vertical="center",
                wrap_text=True,
            )
            cell.border = Border(bottom=ROW_SIDE)
        worksheet.row_dimensions[row_index].height = 23

    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"
    worksheet.print_title_rows = "1:1"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.page_margins = PageMargins(
        left=0.25,
        right=0.25,
        top=0.35,
        bottom=0.35,
        header=0.10,
        footer=0.10,
    )
    widths = (30, 9, 20, 31, 12, 18, 14)
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    output_path = TABLE_DIR / "TableS6_sex_adjusted_biomarkers.xlsx"
    workbook.save(output_path)
    print(f"[PASS] article table written: {output_path}")
    return output_path


def choose_diagnostic_specs(columns: list[str] | None) -> tuple[AncovaSpec, ...]:
    available = {spec.column: spec for spec in ANCOVA_SPECS}
    if columns is None:
        return ANCOVA_SPECS

    unknown = sorted(set(columns).difference(available))
    if unknown:
        raise ValueError(
            f"Unknown diagnostic variable(s): {unknown}. "
            f"Valid columns are: {sorted(available)}"
        )
    ordered_unique = list(dict.fromkeys(columns))
    return tuple(available[column] for column in ordered_unique)


def residual_summary(model: RegressionResultsWrapper) -> dict[str, float]:
    residuals = np.asarray(model.resid, dtype=float)
    shapiro_statistic, shapiro_p = stats.shapiro(residuals)
    return {
        "n": float(residuals.size),
        "shapiro_statistic": float(shapiro_statistic),
        "shapiro_p": float(shapiro_p),
        "skewness": float(stats.skew(residuals, bias=False)),
        "excess_kurtosis": float(stats.kurtosis(residuals, fisher=True, bias=False)),
    }


def print_residual_summary(
    spec: AncovaSpec,
    original_model: RegressionResultsWrapper,
    log_model: RegressionResultsWrapper,
) -> None:
    original = residual_summary(original_model)
    logged = residual_summary(log_model)
    print(f"\n{spec.label} [{spec.column}] — selected scale: {spec.scale}")
    print(
        "  Original scale: "
        f"n={int(original['n'])}, Shapiro P={fmt_p(original['shapiro_p'])}, "
        f"skewness={original['skewness']:.3f}, "
        f"excess kurtosis={original['excess_kurtosis']:.3f}"
    )
    print(
        "  Log scale:      "
        f"n={int(logged['n'])}, Shapiro P={fmt_p(logged['shapiro_p'])}, "
        f"skewness={logged['skewness']:.3f}, "
        f"excess kurtosis={logged['excess_kurtosis']:.3f}"
    )


def plot_model_row(
    axes: np.ndarray,
    model: RegressionResultsWrapper,
    row_index: int,
    scale_label: str,
    selected: bool,
) -> None:
    residuals = np.asarray(model.resid, dtype=float)
    fitted = np.asarray(model.fittedvalues, dtype=float)
    influence = model.get_influence()
    studentized = np.asarray(influence.resid_studentized_internal, dtype=float)
    summary = residual_summary(model)
    selected_text = " [selected]" if selected else ""

    qq_ax = axes[row_index, 0]
    stats.probplot(residuals, dist="norm", plot=qq_ax)
    qq_ax.get_lines()[0].set_markerfacecolor("black")
    qq_ax.get_lines()[0].set_markeredgecolor("black")
    qq_ax.get_lines()[0].set_markersize(4)
    qq_ax.get_lines()[1].set_color("black")
    qq_ax.set_title(
        f"{scale_label}{selected_text}: residual Q-Q\n"
        f"Shapiro P={fmt_p(summary['shapiro_p'])}, skew={summary['skewness']:.2f}"
    )

    residual_ax = axes[row_index, 1]
    residual_ax.scatter(fitted, residuals, s=20, alpha=0.75)
    residual_ax.axhline(0.0, color="black", linewidth=1.0, linestyle="--")
    residual_ax.set_xlabel("Fitted values")
    residual_ax.set_ylabel("Residuals")
    residual_ax.set_title(f"{scale_label}{selected_text}: residuals vs fitted")

    studentized_ax = axes[row_index, 2]
    observation_index = np.arange(1, len(studentized) + 1)
    studentized_ax.scatter(observation_index, studentized, s=20, alpha=0.75)
    studentized_ax.axhline(0.0, color="black", linewidth=1.0)
    studentized_ax.axhline(2.0, color="black", linewidth=0.8, linestyle="--")
    studentized_ax.axhline(-2.0, color="black", linewidth=0.8, linestyle="--")
    studentized_ax.axhline(3.0, color="black", linewidth=0.8, linestyle=":")
    studentized_ax.axhline(-3.0, color="black", linewidth=0.8, linestyle=":")
    studentized_ax.set_xlabel("Complete-case observation index")
    studentized_ax.set_ylabel("Internally studentized residual")
    studentized_ax.set_title(f"{scale_label}{selected_text}: studentized residuals")


def display_residual_diagnostics(df: pd.DataFrame, specs: tuple[AncovaSpec, ...]) -> None:
    import matplotlib as mpl
    import matplotlib.font_manager as font_manager
    import matplotlib.pyplot as plt

    try:
        font_manager.findfont("Arial", fallback_to_default=False)
    except ValueError as exc:
        raise RuntimeError(
            "Required font 'Arial' is not available for diagnostic figures."
        ) from exc
    mpl.rcParams.update({"font.family": "Arial"})

    print("\nResidual diagnostics are display-only; no figure or CSV is saved.")
    print("Close the current figure window to continue to the next biomarker.")

    for spec in specs:
        _, original_model = fit_candidate_model(df, spec, "original")
        _, log_model = fit_candidate_model(df, spec, "log")
        print_residual_summary(spec, original_model, log_model)

        fig, axes = plt.subplots(2, 3, figsize=(13.2, 8.0), constrained_layout=True)
        plot_model_row(
            axes,
            original_model,
            row_index=0,
            scale_label="Original scale",
            selected=spec.scale == "original",
        )
        plot_model_row(
            axes,
            log_model,
            row_index=1,
            scale_label="Log scale",
            selected=spec.scale == "log",
        )
        fig.suptitle(
            f"{spec.label}: ANCOVA residual diagnostics\n"
            "Model for both rows: outcome ~ BE group + sex",
            fontsize=13,
        )
        print(f"Displaying {spec.column}. Close the figure window to continue.")
        plt.show(block=True)
        plt.close(fig)

    print("\nFinished. Residual diagnostic figures were displayed only and were not saved.")


def main() -> None:
    args = parse_args()
    df = load_data()

    results = pd.DataFrame([fit_endpoint(df, spec) for spec in ANCOVA_SPECS])
    results["fdr_bh"] = multipletests(
        results["p_value"].to_numpy(dtype=float), method="fdr_bh"
    )[1]
    formatted = build_formatted_table(results)
    write_table_s6(formatted)

    print("[PASS] sample check: 122 participants; Non-BE=66; BE=56")
    print("[PASS] model unchanged: outcome ~ BE group + sex")
    print("[PASS] HC3 inference and Type III partial eta squared completed")
    print("[PASS] one BH-FDR family applied across eight prespecified biomarkers")
    print("[PASS] module 02 sex-adjusted table workflow completed")

    if args.diagnostic_variables and not args.diagnostics:
        warnings.warn(
            "--diagnostic-variables has no effect unless --diagnostics is also supplied.",
            RuntimeWarning,
            stacklevel=2,
        )
    if args.diagnostics:
        display_residual_diagnostics(df, choose_diagnostic_specs(args.diagnostic_variables))


if __name__ == "__main__":
    main()
