#!/usr/bin/env python3
"""Create article-ready unadjusted preoperative comparison tables.

Input:
    data/preoperative_group_comparison.csv
Outputs:
    outputs/02_preoperative_group_comparison/table/Table1_baseline_characteristics.xlsx
    outputs/02_preoperative_group_comparison/table/Table2_preoperative_biomarkers.xlsx
    outputs/02_preoperative_group_comparison/table/TableS3_unadjusted_baseline_and_weight_outcomes.xlsx
    outputs/02_preoperative_group_comparison/table/TableS4_preoperative_biomarker_differences.xlsx

Endpoint-specific tests, effect directions, bootstrap settings, and domain-wise
FDR families remain fixed to the validated analysis.
"""

from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

import numpy as np
import pandas as pd
from scipy import stats
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from statsmodels.stats.multitest import multipletests
from statsmodels.stats.proportion import (
    confint_proportions_2indep,
)

PROJECT_DIR = Path(__file__).resolve().parents[2]
DATA_FILE = PROJECT_DIR / "data" / "preoperative_group_comparison.csv"
TABLE_DIR = PROJECT_DIR / "outputs" / "02_preoperative_group_comparison" / "table"

N_BOOTSTRAP = 2_000
BOOTSTRAP_SEED = 42
ALPHA = 0.05

TestName = Literal["student", "welch", "mwu", "descriptive_mwu"]


@dataclass(frozen=True)
class VariableSpec:
    column: str
    label: str
    test: TestName
    family: Literal["general", "biomarker"]
    domain: str
    subdomain: str = ""


# The test assigned to each endpoint is the frozen manuscript analysis choice.
# This replication script does not silently change tests when a normality
# p value changes across software implementations or package versions.
GENERAL_SPECS: tuple[VariableSpec, ...] = (
    VariableSpec("age_years", "Age (years)", "student", "general", "Demographics"),
    VariableSpec("bes_score", "BES score", "descriptive_mwu", "general", "Psychometric"),
    VariableSpec("height_cm", "Height (cm)", "student", "general", "Anthropometrics (preoperative)"),
    VariableSpec("baseline_weight_kg", "Weight, preoperative (kg)", "welch", "general", "Anthropometrics (preoperative)"),
    VariableSpec("baseline_bmi_kg_m2", "BMI, preoperative (kg/m²)", "welch", "general", "Anthropometrics (preoperative)"),
    VariableSpec("hospital_stay_days", "Hospital stay (days)", "mwu", "general", "Perioperative details"),
    VariableSpec("operative_time_min", "Operative time (min)", "student", "general", "Perioperative details"),
    VariableSpec("intraoperative_blood_loss_ml", "Intraoperative blood loss (mL)", "mwu", "general", "Perioperative details"),
    VariableSpec("month6_weight_kg", "Weight, 6-month postoperative (kg)", "welch", "general", "Postoperative weight outcomes"),
    VariableSpec("month6_bmi_kg_m2", "BMI, 6-month postoperative (kg/m²)", "student", "general", "Postoperative weight outcomes"),
    VariableSpec("month6_weight_loss_kg", "Weight loss at 6 months (kg)", "student", "general", "Postoperative weight outcomes"),
    VariableSpec("month6_bmi_loss_kg_m2", "BMI loss at 6 months (kg/m²)", "student", "general", "Postoperative weight outcomes"),
    VariableSpec("percent_total_weight_loss", "%TWL", "student", "general", "Postoperative weight outcomes"),
    VariableSpec("percent_excess_weight_loss", "%EWL", "student", "general", "Postoperative weight outcomes"),
)

BIOMARKER_SPECS: tuple[VariableSpec, ...] = (
    VariableSpec("total_protein_g_l", "Total protein (g/L)", "welch", "biomarker", "Nutrition", "Protein nutrition"),
    VariableSpec("calcium_mmol_l", "Ca (mmol/L)", "student", "biomarker", "Nutrition", "Bone mineral nutrition"),
    VariableSpec("vitamin_d_25oh_ng_ml", "25(OH)D (ng/mL)", "mwu", "biomarker", "Nutrition", "Bone mineral nutrition"),
    VariableSpec("iron_umol_l", "Fe (µmol/L)", "student", "biomarker", "Nutrition", "Iron and hematologic nutrition"),
    VariableSpec("hemoglobin_g_l", "Hb (g/L)", "student", "biomarker", "Nutrition", "Iron and hematologic nutrition"),
    VariableSpec("rbc_10e12_l", "RBC (×10¹²/L)", "student", "biomarker", "Nutrition", "Iron and hematologic nutrition"),
    VariableSpec("folate_ng_ml", "Folate (ng/mL)", "mwu", "biomarker", "Nutrition", "Vitamins"),
    VariableSpec("vitamin_b12_pg_ml", "Vitamin B12 (pg/mL)", "mwu", "biomarker", "Nutrition", "Vitamins"),
    VariableSpec("alt_u_l", "ALT (U/L)", "mwu", "biomarker", "Metabolic", "Hepatic metabolism"),
    VariableSpec("ast_u_l", "AST (U/L)", "mwu", "biomarker", "Metabolic", "Hepatic metabolism"),
    VariableSpec("ggt_u_l", "GGT (U/L)", "mwu", "biomarker", "Metabolic", "Hepatic metabolism"),
    VariableSpec("total_cholesterol_mmol_l", "Total cholesterol (mmol/L)", "student", "biomarker", "Metabolic", "Lipid metabolism"),
    VariableSpec("ldl_c_mmol_l", "LDL-C (mmol/L)", "student", "biomarker", "Metabolic", "Lipid metabolism"),
    VariableSpec("hdl_c_mmol_l", "HDL-C (mmol/L)", "student", "biomarker", "Metabolic", "Lipid metabolism"),
    VariableSpec("fasting_plasma_glucose_mmol_l", "FPG (mmol/L)", "mwu", "biomarker", "Metabolic", "Glycemic metabolism"),
    VariableSpec("hba1c_percent", "HbA1c (%)", "mwu", "biomarker", "Metabolic", "Glycemic metabolism"),
    VariableSpec("urea_mmol_l", "Urea (mmol/L)", "student", "biomarker", "Metabolic", "Renal metabolism"),
    VariableSpec("creatinine_umol_l", "Creatinine (µmol/L)", "student", "biomarker", "Metabolic", "Renal metabolism"),
    VariableSpec("uric_acid_umol_l", "Uric acid (µmol/L)", "student", "biomarker", "Metabolic", "Renal metabolism"),
    VariableSpec("sodium_mmol_l", "Na (mmol/L)", "student", "biomarker", "Metabolic", "Electrolyte balance"),
    VariableSpec("potassium_mmol_l", "K (mmol/L)", "student", "biomarker", "Metabolic", "Electrolyte balance"),
    VariableSpec("chloride_mmol_l", "Cl (mmol/L)", "student", "biomarker", "Metabolic", "Electrolyte balance"),
    VariableSpec("wbc_10e9_l", "WBC (×10⁹/L)", "student", "biomarker", "Immune", "White cell profile"),
    VariableSpec("neutrophils_10e9_l", "Neutrophils (×10⁹/L)", "student", "biomarker", "Immune", "White cell profile"),
    VariableSpec("lymphocytes_10e9_l", "Lymphocytes (×10⁹/L)", "student", "biomarker", "Immune", "White cell profile"),
    VariableSpec("monocytes_10e9_l", "Monocytes (×10⁹/L)", "student", "biomarker", "Immune", "White cell profile"),
    VariableSpec("eosinophils_10e9_l", "Eosinophils (×10⁹/L)", "welch", "biomarker", "Immune", "White cell profile"),
    VariableSpec("basophils_10e9_l", "Basophils (×10⁹/L)", "mwu", "biomarker", "Immune", "White cell profile"),
)

ALL_SPECS = GENERAL_SPECS + BIOMARKER_SPECS



def normalise_group(series: pd.Series) -> pd.Series:
    values = series.astype("string").str.strip().str.lower()
    mapped = values.map({"non-be": "Non-BE", "non_be": "Non-BE", "nonbe": "Non-BE", "be": "BE"})
    unknown = sorted(values[mapped.isna() & values.notna()].dropna().unique().tolist())
    if unknown:
        raise ValueError(f"Unrecognised bes_group value(s): {unknown}")
    return mapped.astype("string")


def normalise_sex(series: pd.Series) -> pd.Series:
    values = series.astype("string").str.strip().str.lower()
    mapped = values.map({"female": "Female", "f": "Female", "0": "Female", "male": "Male", "m": "Male", "1": "Male"})
    unknown = sorted(values[mapped.isna() & values.notna()].dropna().unique().tolist())
    if unknown:
        raise ValueError(f"Unrecognised sex value(s): {unknown}")
    return mapped.astype("string")


def load_data() -> pd.DataFrame:
    if not DATA_FILE.exists():
        raise FileNotFoundError(f"Input file not found: {DATA_FILE}")

    df = pd.read_csv(DATA_FILE)
    df.columns = [str(column).strip() for column in df.columns]
    df = df.replace(r"^\s*$", np.nan, regex=True)

    required = {"participant_id", "sex", "bes_group", *(spec.column for spec in ALL_SPECS)}
    missing = sorted(required.difference(df.columns))
    if missing:
        raise ValueError(f"Missing required column(s): {missing}")

    numeric_columns = {"participant_id", *(spec.column for spec in ALL_SPECS)}
    for column in numeric_columns:
        df[column] = pd.to_numeric(df[column], errors="coerce")

    if df["participant_id"].isna().any():
        raise ValueError("participant_id must be complete and numeric.")
    if df["participant_id"].duplicated().any():
        raise ValueError("participant_id must be unique.")

    df["bes_group"] = normalise_group(df["bes_group"])
    df["sex"] = normalise_sex(df["sex"])

    observed_groups = df["bes_group"].value_counts().to_dict()
    if len(df) != 122 or observed_groups != {"Non-BE": 66, "BE": 56}:
        raise AssertionError(
            f"Expected 122 participants (Non-BE=66, BE=56); observed {len(df)} and {observed_groups}."
        )
    observed_sex = pd.crosstab(df["bes_group"], df["sex"]).to_dict()
    expected_sex = {"Female": {"BE": 43, "Non-BE": 37}, "Male": {"BE": 13, "Non-BE": 29}}
    if observed_sex != expected_sex:
        raise AssertionError(f"Unexpected sex-by-group counts: {observed_sex}")

    return df.sort_values("participant_id", kind="mergesort").reset_index(drop=True)


def arrays_by_group(df: pd.DataFrame, column: str) -> tuple[np.ndarray, np.ndarray]:
    nonbe = df.loc[df["bes_group"].eq("Non-BE"), column].dropna().to_numpy(dtype=float)
    be = df.loc[df["bes_group"].eq("BE"), column].dropna().to_numpy(dtype=float)
    if len(nonbe) == 0 or len(be) == 0:
        raise ValueError(f"No usable values in one or both groups for {column}.")
    return nonbe, be


def mean_sd(values: np.ndarray) -> tuple[float, float]:
    return float(np.mean(values)), float(np.std(values, ddof=1))


def median_iqr(values: np.ndarray) -> tuple[float, float, float]:
    q1, median, q3 = np.quantile(values, [0.25, 0.50, 0.75])
    return float(median), float(q1), float(q3)


def hedges_g(nonbe: np.ndarray, be: np.ndarray) -> tuple[float, float, float, float]:
    """Hedges' g and normal-approximation 95% CI, oriented BE - Non-BE."""
    n0, n1 = len(nonbe), len(be)
    s0 = float(np.std(nonbe, ddof=1))
    s1 = float(np.std(be, ddof=1))
    dfree = n0 + n1 - 2
    pooled_variance = ((n0 - 1) * s0**2 + (n1 - 1) * s1**2) / dfree
    pooled_sd = math.sqrt(pooled_variance)
    if not np.isfinite(pooled_sd) or pooled_sd <= 0:
        return np.nan, np.nan, np.nan, np.nan
    cohen_d = (float(np.mean(be)) - float(np.mean(nonbe))) / pooled_sd
    correction = 1.0 - 3.0 / (4.0 * (n0 + n1) - 9.0)
    g = correction * cohen_d
    se_d = math.sqrt((n0 + n1) / (n0 * n1) + cohen_d**2 / (2.0 * dfree))
    se_g = correction * se_d
    z = stats.norm.ppf(1.0 - ALPHA / 2.0)
    return float(g), float(g - z * se_g), float(g + z * se_g), float(cohen_d)


def mean_difference_student(nonbe: np.ndarray, be: np.ndarray) -> tuple[float, float, float]:
    n0, n1 = len(nonbe), len(be)
    s0 = float(np.std(nonbe, ddof=1))
    s1 = float(np.std(be, ddof=1))
    dfree = n0 + n1 - 2
    pooled_variance = ((n0 - 1) * s0**2 + (n1 - 1) * s1**2) / dfree
    standard_error = math.sqrt(pooled_variance * (1.0 / n0 + 1.0 / n1))
    difference = float(np.mean(be) - np.mean(nonbe))
    critical = stats.t.ppf(1.0 - ALPHA / 2.0, dfree)
    return difference, float(difference - critical * standard_error), float(difference + critical * standard_error)


def mean_difference_welch(nonbe: np.ndarray, be: np.ndarray) -> tuple[float, float, float, float]:
    n0, n1 = len(nonbe), len(be)
    s0 = float(np.std(nonbe, ddof=1))
    s1 = float(np.std(be, ddof=1))
    v0, v1 = s0**2 / n0, s1**2 / n1
    standard_error = math.sqrt(v0 + v1)
    dfree = (v0 + v1) ** 2 / (v0**2 / (n0 - 1) + v1**2 / (n1 - 1))
    difference = float(np.mean(be) - np.mean(nonbe))
    critical = stats.t.ppf(1.0 - ALPHA / 2.0, dfree)
    return difference, float(difference - critical * standard_error), float(difference + critical * standard_error), float(dfree)


def hodges_lehmann(nonbe: np.ndarray, be: np.ndarray) -> float:
    return float(np.median(be[:, np.newaxis] - nonbe[np.newaxis, :]))


def cliffs_delta(nonbe: np.ndarray, be: np.ndarray) -> float:
    u_nonbe, _ = stats.mannwhitneyu(nonbe, be, alternative="two-sided", method="auto")
    return float(1.0 - 2.0 * u_nonbe / (len(nonbe) * len(be)))


def bootstrap_two_sample_ci(
    statistic,
    nonbe: np.ndarray,
    be: np.ndarray,
    *,
    seed: int = BOOTSTRAP_SEED,
) -> tuple[float, float]:
    rng = np.random.default_rng(seed)
    n0, n1 = len(nonbe), len(be)
    values = np.empty(N_BOOTSTRAP, dtype=float)
    for index in range(N_BOOTSTRAP):
        sample_nonbe = nonbe[rng.integers(0, n0, n0)]
        sample_be = be[rng.integers(0, n1, n1)]
        values[index] = statistic(sample_nonbe, sample_be)
    lower, upper = np.quantile(values, [ALPHA / 2.0, 1.0 - ALPHA / 2.0])
    return float(lower), float(upper)


def analyse_continuous(spec: VariableSpec, nonbe: np.ndarray, be: np.ndarray) -> dict[str, object]:
    result: dict[str, object] = {
        "column": spec.column,
        "label": spec.label,
        "family": spec.family,
        "domain": spec.domain,
        "subdomain": spec.subdomain,
        "test_code": spec.test,
        "n_nonbe": len(nonbe),
        "n_be": len(be),
        "p_value": np.nan,
        "q_value": np.nan,
        "effect_type": "",
        "effect": np.nan,
        "effect_ci_lower": np.nan,
        "effect_ci_upper": np.nan,
        "difference_type": "",
        "difference": np.nan,
        "difference_ci_lower": np.nan,
        "difference_ci_upper": np.nan,
    }

    if spec.test in {"student", "welch"}:
        mean0, sd0 = mean_sd(nonbe)
        mean1, sd1 = mean_sd(be)
        result.update(
            {
                "summary_type": "mean_sd",
                "nonbe_center": mean0,
                "nonbe_spread_low": sd0,
                "nonbe_spread_high": np.nan,
                "be_center": mean1,
                "be_spread_low": sd1,
                "be_spread_high": np.nan,
            }
        )
        if spec.test == "student":
            _, p_value = stats.ttest_ind(nonbe, be, equal_var=True, alternative="two-sided")
            difference, diff_low, diff_high = mean_difference_student(nonbe, be)
            test_label = "Student's t test"
        else:
            _, p_value = stats.ttest_ind(nonbe, be, equal_var=False, alternative="two-sided")
            difference, diff_low, diff_high, _ = mean_difference_welch(nonbe, be)
            test_label = "Welch's t test"
        g, g_low, g_high, _ = hedges_g(nonbe, be)
        result.update(
            {
                "test_label": test_label,
                "p_value": float(p_value),
                "effect_type": "Hedges' g",
                "effect": g,
                "effect_ci_lower": g_low,
                "effect_ci_upper": g_high,
                "difference_type": "Mean difference",
                "difference": difference,
                "difference_ci_lower": diff_low,
                "difference_ci_upper": diff_high,
            }
        )
        return result

    median0, q10, q30 = median_iqr(nonbe)
    median1, q11, q31 = median_iqr(be)
    hl = hodges_lehmann(nonbe, be)
    hl_low, hl_high = bootstrap_two_sample_ci(hodges_lehmann, nonbe, be)
    result.update(
        {
            "summary_type": "median_iqr",
            "nonbe_center": median0,
            "nonbe_spread_low": q10,
            "nonbe_spread_high": q30,
            "be_center": median1,
            "be_spread_low": q11,
            "be_spread_high": q31,
            "test_label": "Mann–Whitney U test" if spec.test == "mwu" else "Descriptive Hodges–Lehmann contrast",
            "difference_type": "Hodges–Lehmann location-shift estimate",
            "difference": hl,
            "difference_ci_lower": hl_low,
            "difference_ci_upper": hl_high,
        }
    )

    if spec.test == "descriptive_mwu":
        return result

    _, p_value = stats.mannwhitneyu(nonbe, be, alternative="two-sided", method="auto")
    delta = cliffs_delta(nonbe, be)
    delta_low, delta_high = bootstrap_two_sample_ci(cliffs_delta, nonbe, be)
    result.update(
        {
            "p_value": float(p_value),
            "effect_type": "Cliff's delta",
            "effect": delta,
            "effect_ci_lower": delta_low,
            "effect_ci_upper": delta_high,
        }
    )
    return result


def analyse_sex(df: pd.DataFrame) -> dict[str, object]:
    nonbe = df.loc[df["bes_group"].eq("Non-BE"), "sex"]
    be = df.loc[df["bes_group"].eq("BE"), "sex"]
    nonbe_male = int(nonbe.eq("Male").sum())
    be_male = int(be.eq("Male").sum())
    n_nonbe, n_be = len(nonbe), len(be)
    nonbe_female = n_nonbe - nonbe_male
    be_female = n_be - be_male

    contingency = np.array([[be_male, be_female], [nonbe_male, nonbe_female]], dtype=int)
    chi_square, p_value, _, expected = stats.chi2_contingency(contingency, correction=False)
    if float(np.min(expected)) < 5.0:
        _, p_value = stats.fisher_exact(contingency, alternative="two-sided")
        test_label = "Fisher's exact test"
        chi_square = np.nan
    else:
        test_label = "Pearson's chi-square test"

    p_be = be_male / n_be
    p_nonbe = nonbe_male / n_nonbe
    prevalence_ratio = p_be / p_nonbe
    se_log = math.sqrt(1.0 / be_male - 1.0 / n_be + 1.0 / nonbe_male - 1.0 / n_nonbe)
    z = stats.norm.ppf(1.0 - ALPHA / 2.0)
    pr_low = math.exp(math.log(prevalence_ratio) - z * se_log)
    pr_high = math.exp(math.log(prevalence_ratio) + z * se_log)

    risk_difference = p_be - p_nonbe
    rd_low, rd_high = confint_proportions_2indep(
        be_male,
        n_be,
        nonbe_male,
        n_nonbe,
        method="newcomb",
        compare="diff",
        alpha=ALPHA,
        correction=True,
    )

    return {
        "test_label": test_label,
        "chi_square": float(chi_square) if np.isfinite(chi_square) else np.nan,
        "p_value": float(p_value),
        "n_nonbe": n_nonbe,
        "n_be": n_be,
        "nonbe_male": nonbe_male,
        "nonbe_female": nonbe_female,
        "be_male": be_male,
        "be_female": be_female,
        "pr": float(prevalence_ratio),
        "pr_low": float(pr_low),
        "pr_high": float(pr_high),
        "rd": float(risk_difference),
        "rd_low": float(rd_low),
        "rd_high": float(rd_high),
    }



def fmt_number(value: float, digits: int = 2) -> str:
    if not np.isfinite(value):
        return "—"
    rounded = 0.0 if abs(value) < 0.5 * 10 ** (-digits) else value
    return f"{rounded:.{digits}f}"


def fmt_p(value: float) -> str:
    if not np.isfinite(value):
        return "—"
    return "<0.001" if value < 0.001 else f"{value:.3f}"


def fmt_q(value: float) -> str:
    return fmt_p(value)


def fmt_summary(row: pd.Series, group: Literal["nonbe", "be"]) -> str:
    center = float(row[f"{group}_center"])
    low = float(row[f"{group}_spread_low"])
    high = float(row[f"{group}_spread_high"]) if pd.notna(row[f"{group}_spread_high"]) else np.nan
    if row["summary_type"] == "mean_sd":
        return f"{center:.2f} ± {low:.2f}"
    return f"{center:.2f} ({low:.2f}, {high:.2f})"


def fmt_effect(row: pd.Series) -> str:
    if not np.isfinite(float(row["effect"])):
        return "—"
    prefix = "g" if row["effect_type"] == "Hedges' g" else "δ"
    return (
        f"{prefix} {fmt_number(float(row['effect']), 2)} "
        f"({fmt_number(float(row['effect_ci_lower']), 2)}, "
        f"{fmt_number(float(row['effect_ci_upper']), 2)})"
    )


def fmt_difference(row: pd.Series) -> str:
    if not np.isfinite(float(row["difference"])):
        return "—"
    return (
        f"{fmt_number(float(row['difference']), 2)} "
        f"({fmt_number(float(row['difference_ci_lower']), 2)}, "
        f"{fmt_number(float(row['difference_ci_upper']), 2)})"
    )


def build_table_1(results: pd.DataFrame, sex_result: dict[str, object]) -> pd.DataFrame:
    rows: list[dict[str, object]] = [
        {
            "Variable (unit)": "Sex, male",
            "Non-BE n": sex_result["nonbe_male"],
            "Non-BE value": "",
            "BE n": sex_result["be_male"],
            "BE value": "",
            "P value": fmt_p(float(sex_result["p_value"])),
            "Effect size (95% CI)": (
                f"PR {float(sex_result['pr']):.2f} "
                f"({float(sex_result['pr_low']):.2f}, {float(sex_result['pr_high']):.2f})"
            ),
        },
        {
            "Variable (unit)": "Sex, female",
            "Non-BE n": sex_result["nonbe_female"],
            "Non-BE value": "",
            "BE n": sex_result["be_female"],
            "BE value": "",
            "P value": "",
            "Effect size (95% CI)": "",
        },
    ]
    for _, row in results.loc[results["family"].eq("general")].iterrows():
        rows.append(
            {
                "Variable (unit)": row["label"],
                "Non-BE n": int(row["n_nonbe"]),
                "Non-BE value": fmt_summary(row, "nonbe"),
                "BE n": int(row["n_be"]),
                "BE value": fmt_summary(row, "be"),
                "P value": fmt_p(float(row["p_value"])),
                "Effect size (95% CI)": fmt_effect(row),
            }
        )
    return pd.DataFrame(rows)


def build_table_s3(results: pd.DataFrame, sex_result: dict[str, object]) -> pd.DataFrame:
    rows: list[dict[str, object]] = [
        {
            "Section": "Demographics",
            "Indicator (unit)": "Sex",
            "Statistical test": sex_result["test_label"],
            "Difference (95% CI)": (
                f"{fmt_number(float(sex_result['rd']), 2)} "
                f"({fmt_number(float(sex_result['rd_low']), 2)}, "
                f"{fmt_number(float(sex_result['rd_high']), 2)})"
            ),
            "P value": fmt_p(float(sex_result["p_value"])),
        }
    ]
    for _, row in results.loc[results["family"].eq("general")].iterrows():
        rows.append(
            {
                "Section": row["domain"],
                "Indicator (unit)": row["label"],
                "Statistical test": row["test_label"],
                "Difference (95% CI)": fmt_difference(row),
                "P value": fmt_p(float(row["p_value"])),
            }
        )
    return pd.DataFrame(rows)


def build_table_2(results: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for _, row in results.loc[results["family"].eq("biomarker")].iterrows():
        rows.append(
            {
                "Domain": row["domain"],
                "Indicator (unit)": row["label"],
                "Non-BE n": int(row["n_nonbe"]),
                "Non-BE value": fmt_summary(row, "nonbe"),
                "BE n": int(row["n_be"]),
                "BE value": fmt_summary(row, "be"),
                "P value": fmt_p(float(row["p_value"])),
                "Effect size (95% CI)": fmt_effect(row),
            }
        )
    return pd.DataFrame(rows)


def build_table_s4(results: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for _, row in results.loc[results["family"].eq("biomarker")].iterrows():
        rows.append(
            {
                "Domain": row["domain"],
                "Subdomain": row["subdomain"],
                "Indicator (unit)": row["label"],
                "Statistical test": row["test_label"],
                "Difference (95% CI)": fmt_difference(row),
                "P value": fmt_p(float(row["p_value"])),
                "q value (BH-FDR)": fmt_q(float(row["q_value"])),
            }
        )
    return pd.DataFrame(rows)


ARTICLE_FONT = "Arial"
HEADER_SIDE = Side(style="thin", color="808080")
ROW_SIDE = Side(style="hair", color="D9D9D9")
DOMAIN_LABELS = {
    "Nutrition": "Nutrition indicators",
    "Metabolic": "Metabolic indicators",
    "Immune": "Immune indicators",
}


def style_cell(
    cell,
    *,
    bold: bool = False,
    italic: bool = False,
    horizontal: str = "center",
    indent: int = 0,
    wrap_text: bool = False,
    header: bool = False,
) -> None:
    cell.font = Font(name=ARTICLE_FONT, size=9, bold=bold, italic=italic)
    cell.alignment = Alignment(
        horizontal=horizontal,
        vertical="center",
        indent=indent,
        wrap_text=wrap_text,
    )
    cell.border = Border(bottom=HEADER_SIDE if header else ROW_SIDE)


def configure_sheet(
    worksheet,
    *,
    widths: tuple[float, ...],
    freeze_panes: str,
    header_rows: int,
) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = freeze_panes
    worksheet.print_title_rows = f"1:{header_rows}"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.page_margins = PageMargins(
        left=0.25,
        right=0.25,
        top=0.35,
        bottom=0.35,
        header=0.10,
        footer=0.10,
    )
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    for row_index in range(1, header_rows + 1):
        worksheet.row_dimensions[row_index].height = 30 if header_rows == 2 else 34


def write_double_header(
    worksheet,
    first_header: str,
    n_nonbe: int,
    n_be: int,
) -> None:
    worksheet.append(
        [
            first_header,
            f"Non-BE (n={n_nonbe})",
            "",
            f"BE (n={n_be})",
            "",
            "P value",
            "Effect size (95% CI)",
        ]
    )
    worksheet.append(["", "n", "Value", "n", "Value", "", ""])
    for row in worksheet.iter_rows(min_row=1, max_row=2, min_col=1, max_col=7):
        for cell in row:
            style_cell(cell, bold=True, wrap_text=True, header=True)
    for cell_range in ("A1:A2", "B1:C1", "D1:E1", "F1:F2", "G1:G2"):
        worksheet.merge_cells(cell_range)


def append_data_row(worksheet, values: list[object], *, indent: int = 0) -> None:
    worksheet.append(values)
    row_index = worksheet.max_row
    for column_index, cell in enumerate(worksheet[row_index], start=1):
        style_cell(
            cell,
            horizontal="left" if column_index == 1 else "center",
            indent=indent if column_index == 1 else 0,
            wrap_text=True,
        )
    worksheet.row_dimensions[row_index].height = 24


def append_group_row(
    worksheet,
    label: str,
    column_count: int,
    *,
    indent: int = 0,
    italic: bool = False,
) -> None:
    worksheet.append([label, *([""] * (column_count - 1))])
    row_index = worksheet.max_row
    for column_index, cell in enumerate(worksheet[row_index], start=1):
        style_cell(
            cell,
            bold=True,
            italic=italic,
            horizontal="left" if column_index == 1 else "center",
            indent=indent if column_index == 1 else 0,
            wrap_text=True,
        )
    worksheet.row_dimensions[row_index].height = 22


def save_workbook(workbook: Workbook, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    print(f"[PASS] article table written: {output_path}")


def write_table_1(table: pd.DataFrame, n_nonbe: int, n_be: int) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Table 1"
    write_double_header(worksheet, "Variable (unit)", n_nonbe, n_be)
    for row in table.itertuples(index=False, name=None):
        append_data_row(worksheet, list(row))
    configure_sheet(
        worksheet,
        widths=(34, 9, 20, 9, 20, 12, 25),
        freeze_panes="A3",
        header_rows=2,
    )
    save_workbook(workbook, TABLE_DIR / "Table1_baseline_characteristics.xlsx")


def write_table_2(table: pd.DataFrame, n_nonbe: int, n_be: int) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Table 2"
    write_double_header(worksheet, "Indicator (unit)", n_nonbe, n_be)

    current_domain = None
    for row in table.itertuples(index=False):
        if row.Domain != current_domain:
            current_domain = row.Domain
            append_group_row(worksheet, DOMAIN_LABELS[current_domain], 7)
        append_data_row(
            worksheet,
            [
                getattr(row, "_1"),
                row._2,
                row._3,
                row._4,
                row._5,
                row._6,
                row._7,
            ],
            indent=1,
        )
    configure_sheet(
        worksheet,
        widths=(34, 9, 20, 9, 20, 12, 25),
        freeze_panes="A3",
        header_rows=2,
    )
    save_workbook(workbook, TABLE_DIR / "Table2_preoperative_biomarkers.xlsx")


def write_single_header(worksheet, headers: list[str]) -> None:
    worksheet.append(headers)
    for cell in worksheet[1]:
        style_cell(cell, bold=True, wrap_text=True, header=True)


def write_table_s3(table: pd.DataFrame) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Table S3"
    headers = [
        "Indicator (unit)",
        "Statistical test",
        "Difference (95% CI)",
        "P value",
    ]
    write_single_header(worksheet, headers)

    current_section = None
    for row in table.itertuples(index=False):
        if row.Section != current_section:
            current_section = row.Section
            append_group_row(worksheet, current_section, len(headers))
        append_data_row(
            worksheet,
            [row._1, row._2, row._3, row._4],
            indent=1,
        )
    configure_sheet(
        worksheet,
        widths=(38, 28, 28, 12),
        freeze_panes="A2",
        header_rows=1,
    )
    save_workbook(
        workbook,
        TABLE_DIR / "TableS3_unadjusted_baseline_and_weight_outcomes.xlsx",
    )


def write_table_s4(table: pd.DataFrame) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Table S4"
    headers = [
        "Indicator (unit)",
        "Statistical test",
        "Difference (95% CI)",
        "P value",
        "q value (BH-FDR)",
    ]
    write_single_header(worksheet, headers)

    current_domain = None
    current_subdomain = None
    for row in table.itertuples(index=False):
        if row.Domain != current_domain:
            current_domain = row.Domain
            current_subdomain = None
            append_group_row(worksheet, DOMAIN_LABELS[current_domain], len(headers))
        if row.Subdomain != current_subdomain:
            current_subdomain = row.Subdomain
            append_group_row(
                worksheet,
                current_subdomain,
                len(headers),
                indent=1,
                italic=True,
            )
        append_data_row(
            worksheet,
            [row._2, row._3, row._4, row._5, row._6],
            indent=2,
        )
    configure_sheet(
        worksheet,
        widths=(38, 28, 28, 12, 18),
        freeze_panes="A2",
        header_rows=1,
    )
    save_workbook(
        workbook,
        TABLE_DIR / "TableS4_preoperative_biomarker_differences.xlsx",
    )


def main() -> None:
    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    df = load_data()

    result_rows: list[dict[str, object]] = []
    for spec in ALL_SPECS:
        nonbe, be = arrays_by_group(df, spec.column)
        result_rows.append(analyse_continuous(spec, nonbe, be))
    results = pd.DataFrame(result_rows)

    biomarker_mask = results["family"].eq("biomarker")
    for domain in ("Nutrition", "Metabolic", "Immune"):
        mask = biomarker_mask & results["domain"].eq(domain)
        results.loc[mask, "q_value"] = multipletests(
            results.loc[mask, "p_value"].to_numpy(dtype=float),
            method="fdr_bh",
        )[1]

    sex_result = analyse_sex(df)
    table_1 = build_table_1(results, sex_result)
    table_2 = build_table_2(results)
    table_s3 = build_table_s3(results, sex_result)
    table_s4 = build_table_s4(results)

    n_nonbe = int(df["bes_group"].eq("Non-BE").sum())
    n_be = int(df["bes_group"].eq("BE").sum())
    write_table_1(table_1, n_nonbe, n_be)
    write_table_2(table_2, n_nonbe, n_be)
    write_table_s3(table_s3)
    write_table_s4(table_s4)

    print(f"[PASS] sample check: {len(df)} participants; Non-BE={n_nonbe}; BE={n_be}")
    print("[PASS] effect direction: BE minus Non-BE; ratios: BE/Non-BE")
    print("[PASS] endpoint-specific tests retained from the validated analysis")
    print("[PASS] BH-FDR retained within Nutrition, Metabolic, and Immune")
    print(f"[PASS] bootstrap settings retained: {N_BOOTSTRAP} resamples; seed={BOOTSTRAP_SEED}")
    print("[PASS] module 02 unadjusted table workflow completed")


if __name__ == "__main__":
    main()
