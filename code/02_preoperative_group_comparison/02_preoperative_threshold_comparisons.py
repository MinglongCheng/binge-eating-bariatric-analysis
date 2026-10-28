#!/usr/bin/env python3
"""Create article-ready threshold-defined abnormality Table S5.

Input:
    data/preoperative_group_comparison.csv
Output:
    outputs/02_preoperative_group_comparison/table/TableS5_threshold_abnormality.xlsx

The 12 thresholds, contingency tests, PR and RD directions, confidence
intervals, and absence of multiplicity correction remain unchanged.
"""

from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import numpy as np
import pandas as pd
from scipy import stats
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from statsmodels.stats.proportion import (
    confint_proportions_2indep,
)

PROJECT_DIR = Path(__file__).resolve().parents[2]
DATA_FILE = PROJECT_DIR / "data" / "preoperative_group_comparison.csv"
TABLE_DIR = PROJECT_DIR / "outputs" / "02_preoperative_group_comparison" / "table"

ALPHA = 0.05


@dataclass(frozen=True)
class ThresholdSpec:
    value_column: str
    label: str
    rule_text: str
    classifier: Callable[[pd.DataFrame], pd.Series]


def _numeric_rule(df: pd.DataFrame, column: str, predicate) -> pd.Series:
    values = pd.to_numeric(df[column], errors="coerce")
    classified = pd.Series(pd.NA, index=df.index, dtype="Int64")
    observed = values.notna()
    classified.loc[observed] = predicate(values.loc[observed]).astype(int).to_numpy()
    return classified


def _sex_specific_rule(
    df: pd.DataFrame,
    column: str,
    male_predicate,
    female_predicate,
) -> pd.Series:
    values = pd.to_numeric(df[column], errors="coerce")
    sex = df["sex"].astype("string")
    classified = pd.Series(pd.NA, index=df.index, dtype="Int64")
    observed = values.notna() & sex.isin(["Male", "Female"])
    male = observed & sex.eq("Male")
    female = observed & sex.eq("Female")
    classified.loc[male] = male_predicate(values.loc[male]).astype(int).to_numpy()
    classified.loc[female] = female_predicate(values.loc[female]).astype(int).to_numpy()
    return classified


THRESHOLD_SPECS: tuple[ThresholdSpec, ...] = (
    ThresholdSpec(
        "total_protein_g_l",
        "Total protein",
        "abnormal if <68.3 g/L",
        lambda df: _numeric_rule(df, "total_protein_g_l", lambda x: x < 68.3),
    ),
    ThresholdSpec(
        "vitamin_d_25oh_ng_ml",
        "25(OH)D",
        "abnormal if <=30 ng/mL",
        lambda df: _numeric_rule(df, "vitamin_d_25oh_ng_ml", lambda x: x <= 30.0),
    ),
    ThresholdSpec(
        "folate_ng_ml",
        "Folate",
        "abnormal if <6 ng/mL",
        lambda df: _numeric_rule(df, "folate_ng_ml", lambda x: x < 6.0),
    ),
    ThresholdSpec(
        "alt_u_l",
        "ALT",
        "abnormal if >60 U/L in males or >42 U/L in females",
        lambda df: _sex_specific_rule(df, "alt_u_l", lambda x: x > 60.0, lambda x: x > 42.0),
    ),
    ThresholdSpec(
        "ast_u_l",
        "AST",
        "abnormal if >43 U/L",
        lambda df: _numeric_rule(df, "ast_u_l", lambda x: x > 43.0),
    ),
    ThresholdSpec(
        "ggt_u_l",
        "GGT",
        "abnormal if >58 U/L in males or >43 U/L in females",
        lambda df: _sex_specific_rule(df, "ggt_u_l", lambda x: x > 58.0, lambda x: x > 43.0),
    ),
    ThresholdSpec(
        "total_cholesterol_mmol_l",
        "Total cholesterol",
        "abnormal if >=5.2 mmol/L",
        lambda df: _numeric_rule(df, "total_cholesterol_mmol_l", lambda x: x >= 5.2),
    ),
    ThresholdSpec(
        "ldl_c_mmol_l",
        "LDL-C",
        "abnormal if >=3.4 mmol/L",
        lambda df: _numeric_rule(df, "ldl_c_mmol_l", lambda x: x >= 3.4),
    ),
    ThresholdSpec(
        "hdl_c_mmol_l",
        "HDL-C",
        "abnormal if <1.0 mmol/L",
        lambda df: _numeric_rule(df, "hdl_c_mmol_l", lambda x: x < 1.0),
    ),
    ThresholdSpec(
        "fasting_plasma_glucose_mmol_l",
        "FPG",
        "abnormal if >=7.0 mmol/L",
        lambda df: _numeric_rule(df, "fasting_plasma_glucose_mmol_l", lambda x: x >= 7.0),
    ),
    ThresholdSpec(
        "hba1c_percent",
        "HbA1c",
        "abnormal if >=6.5%",
        lambda df: _numeric_rule(df, "hba1c_percent", lambda x: x >= 6.5),
    ),
    ThresholdSpec(
        "uric_acid_umol_l",
        "Uric acid",
        "abnormal if >=420 µmol/L in males or >=360 µmol/L in females",
        lambda df: _sex_specific_rule(
            df,
            "uric_acid_umol_l",
            lambda x: x >= 420.0,
            lambda x: x >= 360.0,
        ),
    ),
)



def normalise_group(series: pd.Series) -> pd.Series:
    values = series.astype("string").str.strip().str.lower()
    mapped = values.map({"non-be": "Non-BE", "non_be": "Non-BE", "nonbe": "Non-BE", "be": "BE"})
    unknown = sorted(values[mapped.isna() & values.notna()].dropna().unique().tolist())
    if unknown:
        raise ValueError(f"Unrecognised bes_group value(s): {unknown}")
    return mapped.astype("string")


def normalise_sex(series: pd.Series) -> pd.Series:
    values = series.astype("string").str.strip().str.lower()
    mapped = values.map({"female": "Female", "f": "Female", "0": "Female", "male": "Male", "m": "Male", "1": "Male"})
    unknown = sorted(values[mapped.isna() & values.notna()].dropna().unique().tolist())
    if unknown:
        raise ValueError(f"Unrecognised sex value(s): {unknown}")
    return mapped.astype("string")


def load_data() -> pd.DataFrame:
    if not DATA_FILE.exists():
        raise FileNotFoundError(f"Input file not found: {DATA_FILE}")
    df = pd.read_csv(DATA_FILE)
    df.columns = [str(column).strip() for column in df.columns]
    df = df.replace(r"^\s*$", np.nan, regex=True)

    required = {"participant_id", "sex", "bes_group", *(spec.value_column for spec in THRESHOLD_SPECS)}
    missing = sorted(required.difference(df.columns))
    if missing:
        raise ValueError(f"Missing required column(s): {missing}")

    df["participant_id"] = pd.to_numeric(df["participant_id"], errors="coerce")
    if df["participant_id"].isna().any() or df["participant_id"].duplicated().any():
        raise ValueError("participant_id must be complete, numeric, and unique.")
    df["bes_group"] = normalise_group(df["bes_group"])
    df["sex"] = normalise_sex(df["sex"])

    for spec in THRESHOLD_SPECS:
        df[spec.value_column] = pd.to_numeric(df[spec.value_column], errors="coerce")

    observed_groups = df["bes_group"].value_counts().to_dict()
    if len(df) != 122 or observed_groups != {"Non-BE": 66, "BE": 56}:
        raise AssertionError(
            f"Expected 122 participants (Non-BE=66, BE=56); observed {len(df)} and {observed_groups}."
        )

    return df.sort_values("participant_id", kind="mergesort").reset_index(drop=True)


def katz_prevalence_ratio(
    be_events: int,
    be_nonevents: int,
    nonbe_events: int,
    nonbe_nonevents: int,
) -> tuple[float, float, float, bool]:
    raw_n_be = be_events + be_nonevents
    raw_n_nonbe = nonbe_events + nonbe_nonevents
    raw_p_be = be_events / raw_n_be
    raw_p_nonbe = nonbe_events / raw_n_nonbe
    raw_pr = raw_p_be / raw_p_nonbe if raw_p_nonbe > 0 else np.nan

    corrected = any(value == 0 for value in (be_events, be_nonevents, nonbe_events, nonbe_nonevents))
    a, b, c, d = map(float, (be_events, be_nonevents, nonbe_events, nonbe_nonevents))
    if corrected:
        a, b, c, d = a + 0.5, b + 0.5, c + 0.5, d + 0.5
    n_be = a + b
    n_nonbe = c + d
    p_be = a / n_be
    p_nonbe = c / n_nonbe
    if p_be <= 0 or p_nonbe <= 0:
        return float(raw_pr), np.nan, np.nan, corrected
    standard_error = math.sqrt(1.0 / a - 1.0 / n_be + 1.0 / c - 1.0 / n_nonbe)
    z = stats.norm.ppf(1.0 - ALPHA / 2.0)
    ratio_for_ci = p_be / p_nonbe
    lower = math.exp(math.log(ratio_for_ci) - z * standard_error)
    upper = math.exp(math.log(ratio_for_ci) + z * standard_error)
    return float(raw_pr), float(lower), float(upper), corrected


def analyse_endpoint(df: pd.DataFrame, spec: ThresholdSpec) -> dict[str, object]:
    derived = spec.classifier(df)
    complete = pd.DataFrame({"bes_group": df["bes_group"], "abnormal": derived}).dropna().copy()
    complete["abnormal"] = complete["abnormal"].astype(int)

    counts: dict[str, dict[str, int]] = {}
    for group in ("Non-BE", "BE"):
        values = complete.loc[complete["bes_group"].eq(group), "abnormal"]
        normal = int(values.eq(0).sum())
        abnormal = int(values.eq(1).sum())
        counts[group] = {"normal": normal, "abnormal": abnormal, "total": normal + abnormal}

    be_abnormal = counts["BE"]["abnormal"]
    be_normal = counts["BE"]["normal"]
    nonbe_abnormal = counts["Non-BE"]["abnormal"]
    nonbe_normal = counts["Non-BE"]["normal"]
    contingency = np.array(
        [[be_abnormal, be_normal], [nonbe_abnormal, nonbe_normal]],
        dtype=int,
    )

    chi_square, p_chi, _, expected = stats.chi2_contingency(contingency, correction=False)
    expected_min = float(np.min(expected))
    if expected_min < 5.0 or bool((contingency == 0).any()):
        _, p_value = stats.fisher_exact(contingency, alternative="two-sided")
        test_label = "Fisher's exact test"
        reported_chi_square = np.nan
    else:
        p_value = p_chi
        test_label = "Pearson's chi-square test"
        reported_chi_square = float(chi_square)

    pr, pr_low, pr_high, pr_corrected = katz_prevalence_ratio(
        be_abnormal,
        be_normal,
        nonbe_abnormal,
        nonbe_normal,
    )

    n_be = counts["BE"]["total"]
    n_nonbe = counts["Non-BE"]["total"]
    p_be = be_abnormal / n_be
    p_nonbe = nonbe_abnormal / n_nonbe
    risk_difference = p_be - p_nonbe
    rd_low, rd_high = confint_proportions_2indep(
        be_abnormal,
        n_be,
        nonbe_abnormal,
        n_nonbe,
        method="newcomb",
        compare="diff",
        alpha=ALPHA,
    )


    return {
        "indicator": spec.label,
        "value_column": spec.value_column,
        "threshold_rule": spec.rule_text,
        "n_nonbe": n_nonbe,
        "nonbe_normal": nonbe_normal,
        "nonbe_abnormal": nonbe_abnormal,
        "n_be": n_be,
        "be_normal": be_normal,
        "be_abnormal": be_abnormal,
        "test": test_label,
        "chi_square": reported_chi_square,
        "expected_min": expected_min,
        "p_value": float(p_value),
        "prevalence_ratio": pr,
        "pr_ci_lower": pr_low,
        "pr_ci_upper": pr_high,
        "pr_continuity_correction_used": pr_corrected,
        "risk_difference": float(risk_difference),
        "rd_ci_lower": float(rd_low),
        "rd_ci_upper": float(rd_high),
    }


def fmt_p(value: float) -> str:
    if not np.isfinite(value):
        return "—"
    return "<0.001" if value < 0.001 else f"{value:.3f}"


def fmt_number(value: float, digits: int = 2) -> str:
    if not np.isfinite(value):
        return "—"
    rounded = 0.0 if abs(value) < 0.5 * 10 ** (-digits) else value
    return f"{rounded:.{digits}f}"


def fmt_count_percent(count: int, total: int) -> str:
    return f"{count} ({100.0 * count / total:.1f}%)"


def build_formatted_table(results: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for _, result in results.iterrows():
        rows.append(
            {
                "Indicator / status": result["indicator"],
                "Non-BE n (%)": int(result["n_nonbe"]),
                "BE n (%)": int(result["n_be"]),
                "Test": result["test"],
                "Statistic": fmt_number(float(result["chi_square"]), 3),
                "P value": fmt_p(float(result["p_value"])),
                "PR (95% CI)": (
                    f"{fmt_number(float(result['prevalence_ratio']), 2)} "
                    f"({fmt_number(float(result['pr_ci_lower']), 2)}, "
                    f"{fmt_number(float(result['pr_ci_upper']), 2)})"
                ),
                "RD (95% CI)": (
                    f"{fmt_number(float(result['risk_difference']), 2)} "
                    f"({fmt_number(float(result['rd_ci_lower']), 2)}, "
                    f"{fmt_number(float(result['rd_ci_upper']), 2)})"
                ),
            }
        )
        rows.append(
            {
                "Indicator / status": "Normal",
                "Non-BE n (%)": fmt_count_percent(int(result["nonbe_normal"]), int(result["n_nonbe"])),
                "BE n (%)": fmt_count_percent(int(result["be_normal"]), int(result["n_be"])),
                "Test": "",
                "Statistic": "",
                "P value": "",
                "PR (95% CI)": "",
                "RD (95% CI)": "",
            }
        )
        rows.append(
            {
                "Indicator / status": "Abnormal",
                "Non-BE n (%)": fmt_count_percent(int(result["nonbe_abnormal"]), int(result["n_nonbe"])),
                "BE n (%)": fmt_count_percent(int(result["be_abnormal"]), int(result["n_be"])),
                "Test": "",
                "Statistic": "",
                "P value": "",
                "PR (95% CI)": "",
                "RD (95% CI)": "",
            }
        )
    return pd.DataFrame(rows)


ARTICLE_FONT = "Arial"
HEADER_SIDE = Side(style="thin", color="808080")
ROW_SIDE = Side(style="hair", color="D9D9D9")


def style_cell(
    cell,
    *,
    bold: bool = False,
    horizontal: str = "center",
    indent: int = 0,
    header: bool = False,
) -> None:
    cell.font = Font(name=ARTICLE_FONT, size=9, bold=bold)
    cell.alignment = Alignment(
        horizontal=horizontal,
        vertical="center",
        indent=indent,
        wrap_text=True,
    )
    cell.border = Border(bottom=HEADER_SIDE if header else ROW_SIDE)


def write_table_s5(table: pd.DataFrame) -> Path:
    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Table S5"
    worksheet.append(list(table.columns))
    for cell in worksheet[1]:
        style_cell(cell, bold=True, header=True)
    worksheet.row_dimensions[1].height = 38

    for row in table.itertuples(index=False, name=None):
        worksheet.append(list(row))
        row_index = worksheet.max_row
        is_status = row[0] in {"Normal", "Abnormal"}
        for column_index, cell in enumerate(worksheet[row_index], start=1):
            style_cell(
                cell,
                bold=not is_status,
                horizontal="left" if column_index == 1 else "center",
                indent=1 if is_status and column_index == 1 else 0,
            )
        worksheet.row_dimensions[row_index].height = 22

    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"
    worksheet.print_title_rows = "1:1"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.page_margins = PageMargins(
        left=0.20,
        right=0.20,
        top=0.30,
        bottom=0.30,
        header=0.10,
        footer=0.10,
    )
    widths = (32, 16, 16, 24, 12, 11, 21, 23)
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    output_path = TABLE_DIR / "TableS5_threshold_abnormality.xlsx"
    workbook.save(output_path)
    print(f"[PASS] article table written: {output_path}")
    return output_path


def main() -> None:
    df = load_data()
    results = pd.DataFrame([analyse_endpoint(df, spec) for spec in THRESHOLD_SPECS])
    formatted = build_formatted_table(results)
    write_table_s5(formatted)

    print("[PASS] sample check: 122 participants; Non-BE=66; BE=56")
    print("[PASS] 12 prespecified abnormality definitions retained")
    print("[PASS] PR direction: BE/Non-BE; RD direction: BE minus Non-BE")
    print("[PASS] threshold tests and confidence intervals unchanged")
    print("[PASS] module 02 threshold table workflow completed")


if __name__ == "__main__":
    main()
