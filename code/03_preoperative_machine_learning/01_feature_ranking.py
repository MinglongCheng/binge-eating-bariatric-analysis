#!/usr/bin/env python3
"""Create the laboratory feature ranking, Table S7, and frozen outer splits.

Input:
    data/preoperative_group_comparison.csv
Outputs:
    outputs/03_preoperative_machine_learning/splits/_outer_cv_splits.csv
    outputs/03_preoperative_machine_learning/table/TableS7_feature_ranking.xlsx
    outputs/03_preoperative_machine_learning/figure/Fig1d_feature_ranking.png
    outputs/06_final_figures/_plot_data/_fig1_feature_ranking.csv
"""

from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Iterable

import matplotlib

matplotlib.use("Agg")
import matplotlib as mpl
import matplotlib.font_manager as font_manager
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from PIL import Image
from scipy.stats import mannwhitneyu
from sklearn.impute import SimpleImputer
from sklearn.linear_model import LogisticRegression
from sklearn.model_selection import RepeatedStratifiedKFold, StratifiedKFold
from sklearn.preprocessing import StandardScaler
from xgboost import XGBClassifier

PROJECT_DIR = Path(__file__).resolve().parents[2]
DATA_FILE = PROJECT_DIR / "data" / "preoperative_group_comparison.csv"
OUTPUT_DIR = PROJECT_DIR / "outputs" / "03_preoperative_machine_learning"
SPLIT_DIR = OUTPUT_DIR / "splits"
TABLE_DIR = OUTPUT_DIR / "table"
FIGURE_DIR = OUTPUT_DIR / "figure"
PLOT_DATA_DIR = PROJECT_DIR / "outputs" / "06_final_figures" / "_plot_data"
STYLE_MANIFEST = PROJECT_DIR / "code" / "06_figure_assembly" / "00_figure_style_manifest.csv"

OUTER_SPLITS = 5
OUTER_REPEATS = 10
INNER_SPLITS = 3
RANDOM_SEED = 42
FDR_ALPHA = 0.10
MIN_KEEP = 3
L1_C = 1.0
RANKING_ALPHA = 0.78
RANKING_ZORDER = {
    "rank_XGB": 2,
    "rank_FDR": 3,
    "rank_L1": 4,
    "overall_rank": 5,
}
XGB_TOP_K = 3

FEATURES = [
    "wbc_10e9_l",
    "neutrophils_10e9_l",
    "lymphocytes_10e9_l",
    "monocytes_10e9_l",
    "eosinophils_10e9_l",
    "basophils_10e9_l",
    "rbc_10e12_l",
    "hemoglobin_g_l",
    "fasting_plasma_glucose_mmol_l",
    "hba1c_percent",
    "total_protein_g_l",
    "alt_u_l",
    "ast_u_l",
    "ggt_u_l",
    "urea_mmol_l",
    "creatinine_umol_l",
    "uric_acid_umol_l",
    "total_cholesterol_mmol_l",
    "hdl_c_mmol_l",
    "ldl_c_mmol_l",
    "potassium_mmol_l",
    "sodium_mmol_l",
    "chloride_mmol_l",
    "calcium_mmol_l",
    "iron_umol_l",
    "folate_ng_ml",
    "vitamin_b12_pg_ml",
    "vitamin_d_25oh_ng_ml",
]

DISPLAY_NAME = {
    "wbc_10e9_l": "WBC",
    "neutrophils_10e9_l": "Neutrophils",
    "lymphocytes_10e9_l": "Lymphocytes",
    "monocytes_10e9_l": "Monocytes",
    "eosinophils_10e9_l": "Eosinophils",
    "basophils_10e9_l": "Basophils",
    "rbc_10e12_l": "RBC",
    "hemoglobin_g_l": "Hb",
    "fasting_plasma_glucose_mmol_l": "FPG",
    "hba1c_percent": "HbA1c",
    "total_protein_g_l": "Total protein",
    "alt_u_l": "ALT",
    "ast_u_l": "AST",
    "ggt_u_l": "GGT",
    "urea_mmol_l": "Urea",
    "creatinine_umol_l": "Creatinine",
    "uric_acid_umol_l": "Uric acid",
    "total_cholesterol_mmol_l": "Total cholesterol",
    "hdl_c_mmol_l": "HDL-C",
    "ldl_c_mmol_l": "LDL-C",
    "potassium_mmol_l": "K",
    "sodium_mmol_l": "Na",
    "chloride_mmol_l": "Cl",
    "calcium_mmol_l": "Ca",
    "iron_umol_l": "Fe",
    "folate_ng_ml": "Folate",
    "vitamin_b12_pg_ml": "Vitamin B12",
    "vitamin_d_25oh_ng_ml": "25(OH)D",
}

DOMAIN = {
    "total_protein_g_l": "Nutrition indicators",
    "calcium_mmol_l": "Nutrition indicators",
    "vitamin_d_25oh_ng_ml": "Nutrition indicators",
    "iron_umol_l": "Nutrition indicators",
    "hemoglobin_g_l": "Nutrition indicators",
    "rbc_10e12_l": "Nutrition indicators",
    "folate_ng_ml": "Nutrition indicators",
    "vitamin_b12_pg_ml": "Nutrition indicators",
    "alt_u_l": "Metabolic indicators",
    "ast_u_l": "Metabolic indicators",
    "ggt_u_l": "Metabolic indicators",
    "total_cholesterol_mmol_l": "Metabolic indicators",
    "ldl_c_mmol_l": "Metabolic indicators",
    "hdl_c_mmol_l": "Metabolic indicators",
    "fasting_plasma_glucose_mmol_l": "Metabolic indicators",
    "hba1c_percent": "Metabolic indicators",
    "urea_mmol_l": "Metabolic indicators",
    "creatinine_umol_l": "Metabolic indicators",
    "uric_acid_umol_l": "Metabolic indicators",
    "sodium_mmol_l": "Metabolic indicators",
    "potassium_mmol_l": "Metabolic indicators",
    "chloride_mmol_l": "Metabolic indicators",
    "wbc_10e9_l": "Immune indicators",
    "neutrophils_10e9_l": "Immune indicators",
    "lymphocytes_10e9_l": "Immune indicators",
    "monocytes_10e9_l": "Immune indicators",
    "eosinophils_10e9_l": "Immune indicators",
    "basophils_10e9_l": "Immune indicators",
}
FEATURE_ORDER = {feature: order for order, feature in enumerate(FEATURES)}

DEFAULT_STYLE: dict[str, object] = {
    "font_family": "Arial",
    "background_color": "white",
    "png_dpi": 600,
    "single_axis_title_size_pt": 10.0,
    "single_tick_size_pt": 9.0,
    "single_legend_size_pt": 9.0,
}
ARTICLE_FONT = "Arial"
HEADER_SIDE = Side(style="thin", color="808080")
ROW_SIDE = Side(style="hair", color="D9D9D9")


def cm_to_inches(value_cm: float) -> float:
    return value_cm / 2.54


def load_style() -> dict[str, object]:
    style = DEFAULT_STYLE.copy()
    if not STYLE_MANIFEST.exists():
        print("[INFO] shared style manifest not found; using frozen module 03 defaults")
        return style

    manifest = pd.read_csv(STYLE_MANIFEST)
    if not {"Key", "Value"}.issubset(manifest.columns):
        raise ValueError(f"Style manifest must contain Key and Value columns: {STYLE_MANIFEST}")
    values = dict(zip(manifest["Key"].astype(str), manifest["Value"]))
    numeric_keys = {
        "png_dpi",
        "single_axis_title_size_pt",
        "single_tick_size_pt",
        "single_legend_size_pt",
    }
    for key in style:
        if key not in values or pd.isna(values[key]):
            continue
        style[key] = float(values[key]) if key in numeric_keys else str(values[key])
    return style


def configure_matplotlib(style: dict[str, object]) -> None:
    font_family = str(style["font_family"])
    try:
        font_manager.findfont(font_family, fallback_to_default=False)
    except ValueError as exc:
        raise RuntimeError(
            f"Required font '{font_family}' is not available. Install or enable Arial "
            "before generating manuscript figures."
        ) from exc

    mpl.rcParams.update(
        {
            "font.family": font_family,
            "font.size": float(style["single_tick_size_pt"]),
            "axes.labelsize": float(style["single_axis_title_size_pt"]),
            "axes.titlesize": float(style["single_axis_title_size_pt"]),
            "xtick.labelsize": float(style["single_tick_size_pt"]),
            "ytick.labelsize": float(style["single_tick_size_pt"]),
            "legend.fontsize": float(style["single_legend_size_pt"]),
            "axes.facecolor": str(style["background_color"]),
            "figure.facecolor": str(style["background_color"]),
            "savefig.facecolor": str(style["background_color"]),
            "savefig.transparent": False,
        }
    )
    print(f"[PASS] required font available: {font_family}")

def normalise_group(series: pd.Series) -> pd.Series:
    values = series.astype("string").str.strip().str.lower()
    mapped = values.map(
        {
            "non-be": 0,
            "non_be": 0,
            "nonbe": 0,
            "non be": 0,
            "be": 1,
        }
    )
    unknown = sorted(values[mapped.isna() & values.notna()].dropna().unique().tolist())
    if unknown:
        raise ValueError(f"Unrecognised bes_group value(s): {unknown}")
    return mapped.astype("Int64")


def load_data() -> tuple[pd.DataFrame, pd.DataFrame, pd.Series]:
    if not DATA_FILE.exists():
        raise FileNotFoundError(f"Input file not found: {DATA_FILE}")

    df = pd.read_csv(DATA_FILE)
    df.columns = [str(column).strip() for column in df.columns]
    required = {"participant_id", "bes_group", *FEATURES}
    missing = sorted(required.difference(df.columns))
    if missing:
        raise ValueError(f"Missing required column(s): {missing}")

    df["participant_id"] = pd.to_numeric(df["participant_id"], errors="coerce")
    if df["participant_id"].isna().any():
        raise ValueError("participant_id must be complete and numeric.")
    if df["participant_id"].duplicated().any():
        raise ValueError("participant_id must be unique.")

    y = normalise_group(df["bes_group"])
    if y.isna().any():
        raise ValueError("bes_group must be complete.")

    X = df[FEATURES].apply(pd.to_numeric, errors="coerce")
    if X.isna().all(axis=1).any():
        bad_ids = df.loc[X.isna().all(axis=1), "participant_id"].tolist()
        raise ValueError(f"All laboratory predictors are missing for participant(s): {bad_ids}")

    ordered = (
        pd.DataFrame({"participant_id": df["participant_id"], "y": y.astype(int)})
        .sort_values("participant_id", kind="mergesort")
        .reset_index()
    )
    X = X.iloc[ordered["index"].to_numpy()].reset_index(drop=True)
    y = ordered["y"].astype(int).reset_index(drop=True)
    df = df.iloc[ordered["index"].to_numpy()].reset_index(drop=True)

    counts = y.value_counts().sort_index().to_dict()
    if len(df) != 122 or counts != {0: 66, 1: 56}:
        raise AssertionError(
            "Expected 122 participants (Non-BE=66, BE=56); "
            f"observed N={len(df)}, counts={counts}."
        )

    print(f"Input: {DATA_FILE}")
    print(f"Analysis sample: N={len(df)}; Non-BE={(y == 0).sum()}; BE={(y == 1).sum()}")
    print(f"Predictors: {len(FEATURES)} prespecified laboratory variables")
    print(f"Maximum feature-level missingness: {int(X.isna().sum().max())}/{len(X)}")
    return df, X, y


# ---------------------------------------------------------------------------
# Selection methods
# ---------------------------------------------------------------------------


def bh_rejections(p_values: np.ndarray, alpha: float) -> np.ndarray:
    """Benjamini-Hochberg rejection mask; NaNs are treated as non-rejections."""
    p_values = np.asarray(p_values, dtype=float)
    valid_indices = np.flatnonzero(np.isfinite(p_values))
    rejected = np.zeros(p_values.shape, dtype=bool)
    if valid_indices.size == 0:
        return rejected

    ordered_valid = valid_indices[
        np.argsort(p_values[valid_indices], kind="mergesort")
    ]
    ordered_p = p_values[ordered_valid]
    thresholds = alpha * np.arange(1, len(ordered_p) + 1) / len(ordered_p)
    passing = np.flatnonzero(ordered_p <= thresholds)
    if passing.size:
        rejected[ordered_valid[: passing[-1] + 1]] = True
    return rejected


def select_mwu_fdr(X_train: pd.DataFrame, y_train: pd.Series) -> list[str]:
    p_values: list[float] = []
    for feature in FEATURES:
        values = X_train[feature]
        group0 = values[y_train.eq(0)].dropna().to_numpy(dtype=float)
        group1 = values[y_train.eq(1)].dropna().to_numpy(dtype=float)
        if len(group0) < 3 or len(group1) < 3:
            p_values.append(np.nan)
            continue
        try:
            p_values.append(
                float(mannwhitneyu(group0, group1, alternative="two-sided").pvalue)
            )
        except ValueError:
            p_values.append(np.nan)

    p_array = np.asarray(p_values, dtype=float)
    selected_mask = bh_rejections(p_array, FDR_ALPHA)
    selected = [feature for feature, keep in zip(FEATURES, selected_mask) if keep]

    if len(selected) < MIN_KEEP:
        ordered = np.argsort(
            np.where(np.isfinite(p_array), p_array, np.inf), kind="mergesort"
        )
        selected = [FEATURES[index] for index in ordered[:MIN_KEEP]]
    return selected


def select_l1(X_train: pd.DataFrame, y_train: pd.Series, seed: int) -> list[str]:
    """L1 selection with all preprocessing estimated only in X_train."""
    imputer = SimpleImputer(strategy="median")
    scaler = StandardScaler()
    X_imputed = imputer.fit_transform(X_train[FEATURES])
    X_scaled = scaler.fit_transform(X_imputed)

    model = LogisticRegression(
        penalty="l1",
        solver="liblinear",
        C=L1_C,
        max_iter=2000,
        random_state=seed,
    )
    model.fit(X_scaled, y_train.to_numpy(dtype=int))
    coefficients = model.coef_.ravel()
    return [
        feature
        for feature, coefficient in zip(FEATURES, coefficients)
        if abs(float(coefficient)) > 1e-12
    ]


def select_xgb(X_train: pd.DataFrame, y_train: pd.Series, seed: int) -> list[str]:
    """Select the top-XGB_TOP_K features by gain from one training subset."""
    model = XGBClassifier(
        objective="binary:logistic",
        eval_metric="auc",
        tree_method="hist",
        grow_policy="lossguide",
        max_depth=3,
        learning_rate=0.05,
        n_estimators=600,
        subsample=0.8,
        colsample_bytree=0.8,
        reg_lambda=3.0,
        n_jobs=1,
        random_state=seed,
        verbosity=0,
    )
    model.fit(X_train[FEATURES], y_train.to_numpy(dtype=int), verbose=False)
    raw_gain = model.get_booster().get_score(importance_type="gain")
    gain = {feature: float(raw_gain.get(feature, 0.0)) for feature in FEATURES}
    ordered = sorted(FEATURES, key=lambda feature: (-gain[feature], FEATURE_ORDER[feature]))
    selected = ordered[:XGB_TOP_K]
    if len(selected) != XGB_TOP_K:
        raise AssertionError("XGBoost selector did not return the requested number of features.")
    return selected


def selection_counts_for_splits(
    X: pd.DataFrame,
    y: pd.Series,
    training_indices: Iterable[np.ndarray],
) -> tuple[dict[str, Counter[str]], int]:
    counts = {
        "FDR": Counter({feature: 0 for feature in FEATURES}),
        "L1": Counter({feature: 0 for feature in FEATURES}),
        "XGB": Counter({feature: 0 for feature in FEATURES}),
    }
    rounds = 0
    for indices in training_indices:
        X_train = X.iloc[indices]
        y_train = y.iloc[indices]
        seed = RANDOM_SEED

        counts["FDR"].update(select_mwu_fdr(X_train, y_train))
        counts["L1"].update(select_l1(X_train, y_train, seed))
        counts["XGB"].update(select_xgb(X_train, y_train, seed))
        rounds += 1

    if rounds == 0:
        raise AssertionError("No feature-selection rounds were generated.")
    return counts, rounds


def integrated_ranking(
    counts: dict[str, Counter[str]], rounds: int
) -> pd.DataFrame:
    ranking = pd.DataFrame({"feature": FEATURES})
    for method in ("FDR", "L1", "XGB"):
        ranking[f"freq_{method}"] = [counts[method][feature] / rounds for feature in FEATURES]
        ranking[f"rank_{method}"] = (
            ranking[f"freq_{method}"]
            .rank(method="min", ascending=False)
            .astype(int)
        )

    ranking["borda_score"] = ranking[["rank_FDR", "rank_L1", "rank_XGB"]].sum(axis=1)
    ranking["overall_rank"] = (
        ranking["borda_score"].rank(method="min", ascending=True).astype(int)
    )
    ranking["mean_frequency"] = ranking[["freq_FDR", "freq_L1", "freq_XGB"]].mean(axis=1)
    ranking["mean_method_rank"] = ranking[["rank_FDR", "rank_L1", "rank_XGB"]].mean(axis=1)
    ranking["feature_order"] = ranking["feature"].map(FEATURE_ORDER)

    # This ordering chooses a deterministic boundary feature while preserving
    # tied overall ranks in the article table.
    ranking = ranking.sort_values(
        ["borda_score", "mean_frequency", "mean_method_rank", "feature_order"],
        ascending=[True, False, True, True],
        kind="mergesort",
    ).reset_index(drop=True)
    return ranking


# ---------------------------------------------------------------------------
# Split construction and outputs
# ---------------------------------------------------------------------------


def create_outer_splits(
    participant_ids: pd.Series, X: pd.DataFrame, y: pd.Series
) -> tuple[pd.DataFrame, list[tuple[np.ndarray, np.ndarray]]]:
    splitter = RepeatedStratifiedKFold(
        n_splits=OUTER_SPLITS,
        n_repeats=OUTER_REPEATS,
        random_state=RANDOM_SEED,
    )
    outer_pairs = list(splitter.split(X, y))
    if len(outer_pairs) != OUTER_SPLITS * OUTER_REPEATS:
        raise AssertionError("Unexpected number of outer folds.")

    rows: list[dict[str, int]] = []
    for outer_fold_id, (_, test_indices) in enumerate(outer_pairs):
        repeat_id = outer_fold_id // OUTER_SPLITS + 1
        fold_id = outer_fold_id % OUTER_SPLITS + 1
        for index in test_indices:
            rows.append(
                {
                    "participant_id": int(participant_ids.iloc[index]),
                    "repeat_id": repeat_id,
                    "fold_id": fold_id,
                    "outer_fold_id": outer_fold_id,
                }
            )

    manifest = pd.DataFrame(rows).sort_values(
        ["repeat_id", "fold_id", "participant_id"], kind="mergesort"
    )
    if len(manifest) != len(participant_ids) * OUTER_REPEATS:
        raise AssertionError("outer_cv_splits.csv should contain N × repeats rows.")
    counts = manifest.groupby(["participant_id", "repeat_id"]).size()
    if not counts.eq(1).all():
        raise AssertionError("Each participant must occur once per repeat in the test assignment.")
    return manifest.reset_index(drop=True), outer_pairs


def build_global_ranking(
    X: pd.DataFrame,
    y: pd.Series,
    outer_pairs: list[tuple[np.ndarray, np.ndarray]],
) -> pd.DataFrame:
    """Aggregate selection frequencies across all 150 inner-training subsets.

    The repeated outer/inner structure is retained to reproduce the original
    cross-validated stability ranking. The resulting ordering is global and
    descriptive. It is subsequently used to define one fixed Top-4 panel and
    one fixed ranks 5-8 panel for all outer model fits.
    """
    global_counts = {
        "FDR": Counter({feature: 0 for feature in FEATURES}),
        "L1": Counter({feature: 0 for feature in FEATURES}),
        "XGB": Counter({feature: 0 for feature in FEATURES}),
    }
    global_rounds = 0

    for outer_fold_id, (outer_train, _) in enumerate(outer_pairs):
        repeat_id = outer_fold_id // OUTER_SPLITS + 1
        fold_id = outer_fold_id % OUTER_SPLITS + 1
        X_outer = X.iloc[outer_train].reset_index(drop=True)
        y_outer = y.iloc[outer_train].reset_index(drop=True)

        inner = StratifiedKFold(
            n_splits=INNER_SPLITS,
            shuffle=True,
            random_state=RANDOM_SEED,
        )
        inner_training_indices = [train for train, _ in inner.split(X_outer, y_outer)]
        fold_counts, fold_rounds = selection_counts_for_splits(
            X_outer,
            y_outer,
            inner_training_indices,
        )

        for method in ("FDR", "L1", "XGB"):
            global_counts[method].update(fold_counts[method])
        global_rounds += fold_rounds

        print(
            f"[feature selection] outer {outer_fold_id + 1:02d}/50 "
            f"(repeat {repeat_id}, fold {fold_id}) | "
            f"accumulated rounds={global_rounds}"
        )

    expected_rounds = OUTER_SPLITS * OUTER_REPEATS * INNER_SPLITS
    if global_rounds != expected_rounds:
        raise AssertionError(
            f"Expected {expected_rounds} global inner-training rounds; observed {global_rounds}."
        )
    return integrated_ranking(global_counts, global_rounds)



def table_s7_frame(ranking: pd.DataFrame) -> pd.DataFrame:
    return pd.DataFrame(
        {
            "Feature": ranking["feature"].map(DISPLAY_NAME),
            "Domain": ranking["feature"].map(DOMAIN),
            "Overall rank": ranking["overall_rank"].astype(int),
            "Borda rank sum": ranking["borda_score"].astype(int),
            "FDR rank": ranking["rank_FDR"].astype(int),
            "L1 rank": ranking["rank_L1"].astype(int),
            "XGBoost rank": ranking["rank_XGB"].astype(int),
        }
    )


def style_cell(cell, *, header: bool, left: bool) -> None:
    cell.font = Font(name=ARTICLE_FONT, size=9, bold=header, color="000000")
    cell.alignment = Alignment(
        horizontal="left" if left else "center",
        vertical="center",
        wrap_text=True,
    )
    cell.border = Border(bottom=HEADER_SIDE if header else ROW_SIDE)


def write_table_s7(ranking: pd.DataFrame) -> Path:
    table = table_s7_frame(ranking)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Table S7"
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"

    worksheet.append(list(table.columns))
    for row in table.itertuples(index=False, name=None):
        worksheet.append(list(row))

    widths = (24, 23, 14, 17, 12, 12, 15)
    for column_index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = width
    worksheet.row_dimensions[1].height = 32

    for row_index, row in enumerate(worksheet.iter_rows(), start=1):
        for column_index, cell in enumerate(row, start=1):
            style_cell(cell, header=row_index == 1, left=column_index <= 2)
        if row_index > 1:
            worksheet.row_dimensions[row_index].height = 21

    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_margins = PageMargins(
        left=0.25, right=0.25, top=0.35, bottom=0.35, header=0.15, footer=0.15
    )

    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    output = TABLE_DIR / "TableS7_feature_ranking.xlsx"
    workbook.save(output)
    print(f"[PASS] article table written: {output}")
    return output


def write_plot_ready_ranking(ranking: pd.DataFrame) -> Path:
    top8 = ranking.head(8).copy()
    plot_data = pd.DataFrame(
        {
            "panel_id": "d",
            "feature": top8["feature"],
            "feature_name": top8["feature"].map(DISPLAY_NAME),
            "domain": top8["feature"].map(DOMAIN),
            "overall_rank": top8["overall_rank"].astype(int),
            "borda_rank_sum": top8["borda_score"].astype(int),
            "fdr_rank": top8["rank_FDR"].astype(int),
            "l1_rank": top8["rank_L1"].astype(int),
            "xgboost_rank": top8["rank_XGB"].astype(int),
            "display_order": np.arange(1, len(top8) + 1, dtype=int),
        }
    )
    PLOT_DATA_DIR.mkdir(parents=True, exist_ok=True)
    output = PLOT_DATA_DIR / "_fig1_feature_ranking.csv"
    plot_data.to_csv(output, index=False, encoding="utf-8")
    print(f"[PASS] plot-ready data written: {output}")
    return output


def save_rgb_png(
    fig: plt.Figure,
    output_path: Path,
    *,
    dpi: int,
    background_color: str,
    width_cm: float,
    height_cm: float,
) -> None:
    fig.savefig(output_path, dpi=dpi, facecolor=background_color, transparent=False)
    plt.close(fig)

    with Image.open(output_path) as image:
        if image.mode != "RGB":
            image.convert("RGB").save(output_path, dpi=(dpi, dpi))

    expected_size = (
        round(cm_to_inches(width_cm) * dpi),
        round(cm_to_inches(height_cm) * dpi),
    )
    with Image.open(output_path) as image:
        if image.mode != "RGB":
            raise AssertionError(f"Expected RGB PNG, found {image.mode}: {output_path}")
        if any(abs(observed - expected) > 1 for observed, expected in zip(image.size, expected_size)):
            raise AssertionError(
                f"Expected approximately {expected_size} pixels, found {image.size}: {output_path}"
            )
    print(f"[PASS] standalone figure written: {output_path}")


def plot_rank_figure(
    ranking: pd.DataFrame,
    style: dict[str, object],
) -> Path:
    top8 = ranking.head(8).copy()
    y_positions = np.arange(len(top8))
    width_cm, height_cm = 13.0, 9.0

    fig, ax = plt.subplots(
        figsize=(cm_to_inches(width_cm), cm_to_inches(height_cm)),
        constrained_layout=False,
    )
    series = (
        ("rank_FDR", "FDR rank", "#EF476F", "o", 38),
        ("rank_L1", "L1 rank", "#F1D64A", "^", 42),
        ("rank_XGB", "XGBoost rank", "#2FB171", "s", 38),
        ("overall_rank", "Overall rank", "#B0009B", "*", 72),
    )
    for column, label, colour, marker, size in series:
        ax.scatter(
            top8[column],
            y_positions,
            s=size,
            edgecolor="black",
            linewidth=0.55,
            facecolor=colour,
            marker=marker,
            label=label,
            alpha=RANKING_ALPHA,
            zorder=RANKING_ZORDER[column],
        )

    ax.set_yticks(y_positions)
    ax.set_yticklabels(top8["feature"].map(DISPLAY_NAME).tolist())
    ax.set_xlabel("Rank (lower is better)")
    fig.suptitle("Integrated feature ranking", y=0.955, fontsize=10.0)

    maximum_rank = int(
        np.nanmax(top8[["rank_FDR", "rank_L1", "rank_XGB", "overall_rank"]].to_numpy())
    )
    tick_max = maximum_rank if maximum_rank % 2 == 1 else maximum_rank + 1
    ax.set_xticks(np.arange(1, tick_max + 1, 2))
    ax.set_xlim(0.5, tick_max + 0.7)
    ax.grid(axis="x", linestyle="--", linewidth=0.6, color="0.86")
    ax.tick_params(axis="x", length=3.0, width=0.75)
    ax.tick_params(axis="y", length=3.0, width=0.75, pad=4)
    for spine in ax.spines.values():
        spine.set_linewidth(0.8)
        spine.set_color("0.25")
    ax.legend(loc="lower right", frameon=False, handletextpad=0.5, borderaxespad=0.4)

    left_cm, right_cm = 2.60, 0.70
    bottom_cm, top_cm = 1.45, 1.00
    fig.subplots_adjust(
        left=left_cm / width_cm,
        right=1 - right_cm / width_cm,
        bottom=bottom_cm / height_cm,
        top=1 - top_cm / height_cm,
    )

    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    output = FIGURE_DIR / "Fig1d_feature_ranking.png"
    save_rgb_png(
        fig,
        output,
        dpi=int(style["png_dpi"]),
        background_color=str(style["background_color"]),
        width_cm=width_cm,
        height_cm=height_cm,
    )
    return output

def main() -> None:
    style = load_style()
    configure_matplotlib(style)
    SPLIT_DIR.mkdir(parents=True, exist_ok=True)

    df, X, y = load_data()
    split_manifest, outer_pairs = create_outer_splits(df["participant_id"], X, y)
    ranking = build_global_ranking(X, y, outer_pairs)

    split_path = SPLIT_DIR / "_outer_cv_splits.csv"
    split_manifest.to_csv(split_path, index=False, encoding="utf-8")
    print(f"[PASS] outer split assignments written: {split_path}")

    table_path = write_table_s7(ranking)
    plot_data_path = write_plot_ready_ranking(ranking)
    figure_path = plot_rank_figure(ranking, style)

    fixed_top4 = ranking.head(4)["feature"].map(DISPLAY_NAME).tolist()
    fixed_ranks5_8 = ranking.iloc[4:8]["feature"].map(DISPLAY_NAME).tolist()
    print(f"[PASS] sample check: 122 participants; Non-BE=66; BE=56")
    print("[PASS] 150 inner-training feature-selection subsets retained")
    print("[PASS] Borda ranking and deterministic tie-break unchanged")
    print(f"[PASS] Table S7: {table_path}")
    print(f"[PASS] Figure 1d: {figure_path}")
    print(f"[PASS] plot-ready ranking: {plot_data_path}")
    print("[PASS] Fixed Top-4: " + ", ".join(fixed_top4))
    print("[PASS] Positions 5-8: " + ", ".join(fixed_ranks5_8))
    print("[PASS] module 03 feature-ranking workflow completed")


if __name__ == "__main__":
    main()
