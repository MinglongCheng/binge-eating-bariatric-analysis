#!/usr/bin/env python3
"""Train repeated nested-CV models and create Table S8.

Run 01_feature_ranking.py first. The analysis methods, fixed feature panels,
outer and inner cross-validation, class weighting, early stopping, native
missing-value handling, bootstrap confidence intervals, and saved model
inventory are unchanged.
"""

from __future__ import annotations

from pathlib import Path
from typing import Callable

import lightgbm as lgb
import numpy as np
import pandas as pd
from catboost import CatBoostClassifier
from lightgbm import LGBMClassifier
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from sklearn.metrics import average_precision_score, brier_score_loss, roc_auc_score
from sklearn.model_selection import ParameterGrid, StratifiedKFold

PROJECT_DIR = Path(__file__).resolve().parents[2]
DATA_FILE = PROJECT_DIR / "data" / "preoperative_group_comparison.csv"
OUTPUT_DIR = PROJECT_DIR / "outputs" / "03_preoperative_machine_learning"
SPLIT_FILE = OUTPUT_DIR / "splits" / "_outer_cv_splits.csv"
TABLE_S7_FILE = OUTPUT_DIR / "table" / "TableS7_feature_ranking.xlsx"
TABLE_DIR = OUTPUT_DIR / "table"
MODEL_DIR = OUTPUT_DIR / "models"
LGB_MODEL_DIR = MODEL_DIR / "lightgbm"
CAT_MODEL_DIR = MODEL_DIR / "catboost"

OUTER_SPLITS = 5
OUTER_REPEATS = 10
INNER_SPLITS = 3
RANDOM_SEED = 42
EARLY_STOPPING_ROUNDS = 50
LGB_MAX_ITERATIONS = 1000
CAT_MAX_ITERATIONS = 2000
N_BOOTSTRAP = 2000
N_THREADS = 1

FEATURES = [
    "wbc_10e9_l",
    "neutrophils_10e9_l",
    "lymphocytes_10e9_l",
    "monocytes_10e9_l",
    "eosinophils_10e9_l",
    "basophils_10e9_l",
    "rbc_10e12_l",
    "hemoglobin_g_l",
    "fasting_plasma_glucose_mmol_l",
    "hba1c_percent",
    "total_protein_g_l",
    "alt_u_l",
    "ast_u_l",
    "ggt_u_l",
    "urea_mmol_l",
    "creatinine_umol_l",
    "uric_acid_umol_l",
    "total_cholesterol_mmol_l",
    "hdl_c_mmol_l",
    "ldl_c_mmol_l",
    "potassium_mmol_l",
    "sodium_mmol_l",
    "chloride_mmol_l",
    "calcium_mmol_l",
    "iron_umol_l",
    "folate_ng_ml",
    "vitamin_b12_pg_ml",
    "vitamin_d_25oh_ng_ml",
]

DISPLAY_NAME = {
    "wbc_10e9_l": "WBC",
    "neutrophils_10e9_l": "Neutrophils",
    "lymphocytes_10e9_l": "Lymphocytes",
    "monocytes_10e9_l": "Monocytes",
    "eosinophils_10e9_l": "Eosinophils",
    "basophils_10e9_l": "Basophils",
    "rbc_10e12_l": "RBC",
    "hemoglobin_g_l": "Hb",
    "fasting_plasma_glucose_mmol_l": "FPG",
    "hba1c_percent": "HbA1c",
    "total_protein_g_l": "Total protein",
    "alt_u_l": "ALT",
    "ast_u_l": "AST",
    "ggt_u_l": "GGT",
    "urea_mmol_l": "Urea",
    "creatinine_umol_l": "Creatinine",
    "uric_acid_umol_l": "Uric acid",
    "total_cholesterol_mmol_l": "Total cholesterol",
    "hdl_c_mmol_l": "HDL-C",
    "ldl_c_mmol_l": "LDL-C",
    "potassium_mmol_l": "K",
    "sodium_mmol_l": "Na",
    "chloride_mmol_l": "Cl",
    "calcium_mmol_l": "Ca",
    "iron_umol_l": "Fe",
    "folate_ng_ml": "Folate",
    "vitamin_b12_pg_ml": "Vitamin B12",
    "vitamin_d_25oh_ng_ml": "25(OH)D",
}
DISPLAY_TO_FEATURE = {display: feature for feature, display in DISPLAY_NAME.items()}

PANEL_LABELS = {
    "all_features": "All features",
    "fixed_top4": "Top-4 (fixed)",
    "fixed_ranks5_8": "Positions 5–8 (fixed)",
}
ALGORITHM_LABELS = {"catboost": "CatBoost", "lightgbm": "LightGBM"}

CAT_GRID = {
    "depth": [3, 4],
    "learning_rate": [0.03, 0.05],
    "l2_leaf_reg": [3.0, 5.0],
}
LGB_GRID = {
    "num_leaves": [15, 31],
    "max_depth": [3, 4],
    "learning_rate": [0.03, 0.05],
    "min_child_samples": [5, 10],
    "subsample": [1.0],
    "colsample_bytree": [1.0],
}

ARTICLE_FONT = "Arial"
HEADER_SIDE = Side(style="thin", color="808080")
ROW_SIDE = Side(style="hair", color="D9D9D9")

def normalise_group(series: pd.Series) -> pd.Series:
    values = series.astype("string").str.strip().str.lower()
    mapped = values.map(
        {"non-be": 0, "non_be": 0, "nonbe": 0, "non be": 0, "be": 1}
    )
    unknown = sorted(values[mapped.isna() & values.notna()].dropna().unique().tolist())
    if unknown:
        raise ValueError(f"Unrecognised bes_group value(s): {unknown}")
    return mapped.astype("Int64")


def load_data() -> tuple[pd.DataFrame, pd.DataFrame, pd.Series]:
    if not DATA_FILE.exists():
        raise FileNotFoundError(f"Input file not found: {DATA_FILE}")
    df = pd.read_csv(DATA_FILE)
    df.columns = [str(column).strip() for column in df.columns]
    required = {"participant_id", "bes_group", *FEATURES}
    missing = sorted(required.difference(df.columns))
    if missing:
        raise ValueError(f"Missing required column(s): {missing}")

    df["participant_id"] = pd.to_numeric(df["participant_id"], errors="coerce")
    if df["participant_id"].isna().any() or df["participant_id"].duplicated().any():
        raise ValueError("participant_id must be complete, numeric, and unique.")
    y = normalise_group(df["bes_group"])
    if y.isna().any():
        raise ValueError("bes_group must be complete.")
    X = df[FEATURES].apply(pd.to_numeric, errors="coerce")

    order = np.argsort(df["participant_id"].to_numpy(), kind="mergesort")
    df = df.iloc[order].reset_index(drop=True)
    X = X.iloc[order].reset_index(drop=True)
    y = y.iloc[order].astype(int).reset_index(drop=True)

    counts = y.value_counts().sort_index().to_dict()
    if len(df) != 122 or counts != {0: 66, 1: 56}:
        raise AssertionError(
            f"Expected N=122 with counts {{0: 66, 1: 56}}; observed N={len(df)}, {counts}."
        )
    return df, X, y


def load_split_manifest(participant_ids: pd.Series) -> pd.DataFrame:
    if not SPLIT_FILE.exists():
        raise FileNotFoundError(
            f"Missing split file: {SPLIT_FILE}\nRun 01_feature_ranking.py first."
        )
    splits = pd.read_csv(SPLIT_FILE)
    required = {"participant_id", "repeat_id", "fold_id", "outer_fold_id"}
    missing = required.difference(splits.columns)
    if missing:
        raise ValueError(f"Split file is missing columns: {sorted(missing)}")

    for column in required:
        splits[column] = pd.to_numeric(splits[column], errors="raise").astype(int)
    if len(splits) != len(participant_ids) * OUTER_REPEATS:
        raise AssertionError("_outer_cv_splits.csv does not contain N × repeats rows.")
    if set(splits["participant_id"]) != set(participant_ids.astype(int)):
        raise AssertionError("Participant IDs in the split file do not match the analysis data.")
    if splits.groupby(["participant_id", "repeat_id"]).size().ne(1).any():
        raise AssertionError("Each participant must have one test-fold assignment per repeat.")
    if sorted(splits["outer_fold_id"].unique()) != list(range(OUTER_SPLITS * OUTER_REPEATS)):
        raise AssertionError("outer_fold_id must run from 0 to 49.")
    return splits


def build_fold_assignment_matrix(
    participant_ids: pd.Series, splits: pd.DataFrame
) -> np.ndarray:
    """Return the held-out fold (1-5) for every participant and repeat.

    Rows correspond to repeats and columns to the participant order used by the
    analysis data. The matrix is used only for Table S8 aggregation and
    bootstrap resampling; it does not create or alter any split.
    """
    id_to_index = {int(pid): index for index, pid in enumerate(participant_ids)}
    fold_matrix = np.full(
        (OUTER_REPEATS, len(participant_ids)), -1, dtype=np.int8
    )
    for row in splits.itertuples(index=False):
        participant_id = int(row.participant_id)
        repeat_index = int(row.repeat_id) - 1
        if participant_id not in id_to_index:
            raise AssertionError(
                f"Split participant {participant_id} is absent from analysis data."
            )
        fold_matrix[repeat_index, id_to_index[participant_id]] = int(row.fold_id)

    if np.any(fold_matrix < 1) or np.any(fold_matrix > OUTER_SPLITS):
        raise AssertionError("Every participant/repeat must map to one fold from 1 to 5.")
    return fold_matrix


def load_fixed_feature_panels() -> dict[str, list[str]]:
    """Read the deterministic final ordering from Table S7."""
    if not TABLE_S7_FILE.exists():
        raise FileNotFoundError(
            f"Missing Table S7: {TABLE_S7_FILE}\nRun 01_feature_ranking.py first."
        )
    table = pd.read_excel(TABLE_S7_FILE, sheet_name="Table S7")
    if "Feature" not in table.columns or len(table) < 8:
        raise ValueError("Table S7 must contain at least eight ranked features.")

    ordered: list[str] = []
    for display in table["Feature"].iloc[:8].astype(str):
        if display not in DISPLAY_TO_FEATURE:
            raise ValueError(f"Cannot map Table S7 feature name: {display!r}")
        ordered.append(DISPLAY_TO_FEATURE[display])

    if len(ordered) != 8 or len(set(ordered)) != 8:
        raise AssertionError("The first eight Table S7 rows must be eight unique features.")

    panels = {
        "all_features": FEATURES.copy(),
        "fixed_top4": ordered[:4],
        "fixed_ranks5_8": ordered[4:8],
    }
    if set(panels["fixed_top4"]) & set(panels["fixed_ranks5_8"]):
        raise AssertionError("Fixed Top-4 and fixed ranks 5-8 must be disjoint.")
    return panels


def outer_indices(
    df: pd.DataFrame, splits: pd.DataFrame, outer_fold_id: int
) -> tuple[np.ndarray, np.ndarray, int, int]:
    fold_rows = splits.loc[splits["outer_fold_id"].eq(outer_fold_id)]
    if fold_rows.empty:
        raise AssertionError(f"No test assignments found for outer fold {outer_fold_id}.")
    repeat_values = fold_rows["repeat_id"].unique()
    fold_values = fold_rows["fold_id"].unique()
    if len(repeat_values) != 1 or len(fold_values) != 1:
        raise AssertionError("Each outer_fold_id must map to one repeat and one fold.")
    repeat_id, fold_id = int(repeat_values[0]), int(fold_values[0])
    test_ids = set(fold_rows["participant_id"].astype(int))
    test_mask = df["participant_id"].astype(int).isin(test_ids).to_numpy()
    test_indices = np.flatnonzero(test_mask)
    train_indices = np.flatnonzero(~test_mask)
    if len(test_indices) not in {24, 25} or len(train_indices) + len(test_indices) != len(df):
        raise AssertionError("Unexpected outer train/test sizes.")
    return train_indices, test_indices, repeat_id, fold_id


# ---------------------------------------------------------------------------
# Model construction, tuning, and prediction
# ---------------------------------------------------------------------------


def class_weights(y_train: pd.Series) -> tuple[list[float], float]:
    y_array = y_train.to_numpy(dtype=int)
    n0 = int(np.sum(y_array == 0))
    n1 = int(np.sum(y_array == 1))
    if n0 == 0 or n1 == 0:
        raise ValueError("Both classes must be present in every training subset.")
    total = n0 + n1
    cat_weights = [total / (2.0 * n0), total / (2.0 * n1)]
    scale_pos_weight = n0 / n1
    return cat_weights, float(scale_pos_weight)


def build_lightgbm(
    params: dict[str, object],
    y_train: pd.Series,
    seed: int,
    n_estimators: int,
) -> LGBMClassifier:
    _, scale_pos_weight = class_weights(y_train)
    return LGBMClassifier(
        objective="binary",
        metric="auc",
        n_estimators=int(n_estimators),
        random_state=int(seed),
        n_jobs=N_THREADS,
        scale_pos_weight=scale_pos_weight,
        deterministic=True,
        force_col_wise=True,
        verbosity=-1,
        **params,
    )


def build_catboost(
    params: dict[str, object],
    y_train: pd.Series,
    seed: int,
    iterations: int,
    early_stopping: bool,
) -> CatBoostClassifier:
    cat_weights, _ = class_weights(y_train)
    settings: dict[str, object] = {
        "loss_function": "Logloss",
        "eval_metric": "AUC",
        "iterations": int(iterations),
        "random_seed": int(seed),
        "thread_count": N_THREADS,
        "class_weights": cat_weights,
        "allow_writing_files": False,
        "verbose": False,
        **params,
    }
    if early_stopping:
        settings.update(
            {
                "od_type": "Iter",
                "od_wait": EARLY_STOPPING_ROUNDS,
                "use_best_model": True,
            }
        )
    else:
        settings["use_best_model"] = False
    return CatBoostClassifier(**settings)


def parameter_complexity(algo: str, params: dict[str, object]) -> tuple[float, ...]:
    """Prospective deterministic tie-break favouring simpler/regularized models."""
    if algo == "catboost":
        return (
            float(params["depth"]),
            -float(params["l2_leaf_reg"]),
            float(params["learning_rate"]),
        )
    if algo == "lightgbm":
        return (
            float(params["max_depth"]),
            float(params["num_leaves"]),
            -float(params["min_child_samples"]),
            float(params["learning_rate"]),
        )
    raise ValueError(algo)


def tune_and_refit(
    algo: str,
    X_train: pd.DataFrame,
    y_train: pd.Series,
    X_test: pd.DataFrame,
    features: list[str],
    outer_fold_id: int,
) -> tuple[np.ndarray, object, dict[str, object], int, float]:
    grid = CAT_GRID if algo == "catboost" else LGB_GRID
    inner = StratifiedKFold(
        n_splits=INNER_SPLITS,
        shuffle=True,
        random_state=RANDOM_SEED + outer_fold_id,
    )
    inner_pairs = list(inner.split(X_train, y_train))

    best_record: tuple[float, float, tuple[float, ...], dict[str, object], list[int]] | None = None

    for params in ParameterGrid(grid):
        inner_auc: list[float] = []
        best_iterations: list[int] = []

        for inner_fold_id, (inner_train, inner_valid) in enumerate(inner_pairs):
            X_inner_train = X_train.iloc[inner_train][features]
            y_inner_train = y_train.iloc[inner_train]
            X_inner_valid = X_train.iloc[inner_valid][features]
            y_inner_valid = y_train.iloc[inner_valid]
            seed = RANDOM_SEED + outer_fold_id * 100 + inner_fold_id

            if algo == "lightgbm":
                model = build_lightgbm(params, y_inner_train, seed, LGB_MAX_ITERATIONS)
                model.fit(
                    X_inner_train,
                    y_inner_train,
                    eval_set=[(X_inner_valid, y_inner_valid)],
                    eval_metric="auc",
                    callbacks=[
                        lgb.early_stopping(EARLY_STOPPING_ROUNDS, verbose=False),
                        lgb.log_evaluation(period=0),
                    ],
                )
                best_iteration = int(getattr(model, "best_iteration_", 0) or LGB_MAX_ITERATIONS)
                probability = model.predict_proba(X_inner_valid)[:, 1]
            elif algo == "catboost":
                model = build_catboost(
                    params,
                    y_inner_train,
                    seed,
                    CAT_MAX_ITERATIONS,
                    early_stopping=True,
                )
                model.fit(
                    X_inner_train,
                    y_inner_train,
                    eval_set=(X_inner_valid, y_inner_valid),
                    verbose=False,
                )
                best_index = model.get_best_iteration()
                best_iteration = (
                    CAT_MAX_ITERATIONS
                    if best_index is None or int(best_index) < 0
                    else int(best_index) + 1
                )
                probability = model.predict_proba(X_inner_valid)[:, 1]
            else:
                raise ValueError(algo)

            inner_auc.append(float(roc_auc_score(y_inner_valid, probability)))
            best_iterations.append(max(1, best_iteration))

        mean_auc = float(np.mean(inner_auc))
        sd_auc = float(np.std(inner_auc, ddof=1))
        complexity = parameter_complexity(algo, params)
        record = (mean_auc, sd_auc, complexity, dict(params), best_iterations)

        if best_record is None:
            best_record = record
        else:
            current_key = (-record[0], record[1], record[2])
            best_key = (-best_record[0], best_record[1], best_record[2])
            if current_key < best_key:
                best_record = record

    if best_record is None:
        raise RuntimeError("No hyperparameter candidate was evaluated.")

    best_auc, _, _, best_params, iteration_values = best_record
    final_iterations = max(1, int(round(float(np.median(iteration_values)))))
    final_seed = RANDOM_SEED + outer_fold_id

    if algo == "lightgbm":
        final_model = build_lightgbm(best_params, y_train, final_seed, final_iterations)
        final_model.fit(X_train[features], y_train)
        test_probability = final_model.predict_proba(X_test[features])[:, 1]
    else:
        final_model = build_catboost(
            best_params,
            y_train,
            final_seed,
            final_iterations,
            early_stopping=False,
        )
        final_model.fit(X_train[features], y_train, verbose=False)
        test_probability = final_model.predict_proba(X_test[features])[:, 1]

    return test_probability, final_model, best_params, final_iterations, best_auc


def save_interpretation_model(
    panel: str,
    algo: str,
    model: object,
    repeat_id: int,
    fold_id: int,
) -> Path:
    """Save models required by 03_model_interpretation.py."""
    if panel not in {"all_features", "fixed_top4"}:
        raise ValueError(f"Panel {panel!r} is not retained for interpretation.")
    if algo == "lightgbm":
        path = LGB_MODEL_DIR / (
            f"{panel}_lightgbm_repeat{repeat_id:02d}_fold{fold_id:02d}.txt"
        )
        if not isinstance(model, LGBMClassifier):
            raise TypeError("Expected an LGBMClassifier.")
        model.booster_.save_model(str(path))
        return path
    if algo == "catboost":
        path = CAT_MODEL_DIR / (
            f"{panel}_catboost_repeat{repeat_id:02d}_fold{fold_id:02d}.cbm"
        )
        if not isinstance(model, CatBoostClassifier):
            raise TypeError("Expected a CatBoostClassifier.")
        model.save_model(str(path), format="cbm")
        return path
    raise ValueError(algo)


# ---------------------------------------------------------------------------
# Table S8 summaries
# ---------------------------------------------------------------------------


def metric_function(name: str) -> Callable[[np.ndarray, np.ndarray], float]:
    if name == "AUPRC":
        return lambda y, p: float(average_precision_score(y, p))
    if name == "AUROC":
        return lambda y, p: float(roc_auc_score(y, p))
    if name == "Brier score":
        return lambda y, p: float(brier_score_loss(y, p))
    raise ValueError(name)


def repeat_level_fold_scores(
    y: np.ndarray,
    prediction_by_repeat: np.ndarray,
    fold_by_repeat: np.ndarray,
    name: str,
) -> np.ndarray:
    """Compute one fold-first performance estimate for each CV repeat.

    AUROC and AUPRC are nonlinear, non-decomposable metrics. They are therefore
    calculated separately in each held-out fold and then averaged within the
    repeat. Brier loss is decomposable, so its fold values are weighted by the
    corresponding held-out sample sizes.
    """
    metric = metric_function(name)
    repeat_values = np.empty(OUTER_REPEATS, dtype=float)

    for repeat in range(OUTER_REPEATS):
        fold_values: list[float] = []
        fold_sizes: list[int] = []
        for fold_id in range(1, OUTER_SPLITS + 1):
            within_fold = fold_by_repeat[repeat] == fold_id
            y_fold = y[within_fold]
            probability_fold = prediction_by_repeat[repeat, within_fold]
            if name in {"AUPRC", "AUROC"} and np.unique(y_fold).size < 2:
                raise ValueError("Each outer-test fold must contain both classes.")
            fold_values.append(metric(y_fold, probability_fold))
            fold_sizes.append(int(np.sum(within_fold)))

        if name == "Brier score":
            repeat_values[repeat] = float(
                np.average(fold_values, weights=fold_sizes)
            )
        else:
            repeat_values[repeat] = float(np.mean(fold_values))

    return repeat_values


def make_stratified_bootstrap_weights(
    y: np.ndarray,
    fold_by_repeat: np.ndarray,
    rng: np.random.Generator,
) -> np.ndarray:
    """Generate valid participant-level stratified bootstrap multiplicities.

    Each row contains the number of times every participant is selected. The
    same bootstrap weights are reused for all panels, algorithms, and metrics.
    Rare draws that remove either class from any held-out fold are discarded.
    """
    class0 = np.flatnonzero(y == 0)
    class1 = np.flatnonzero(y == 1)
    probability0 = np.full(len(class0), 1.0 / len(class0))
    probability1 = np.full(len(class1), 1.0 / len(class1))
    accepted: list[np.ndarray] = []
    n_accepted = 0

    while n_accepted < N_BOOTSTRAP:
        batch_size = max(256, N_BOOTSTRAP - n_accepted)
        weights = np.zeros((batch_size, len(y)), dtype=np.int16)
        weights[:, class0] = rng.multinomial(
            len(class0), probability0, size=batch_size
        )
        weights[:, class1] = rng.multinomial(
            len(class1), probability1, size=batch_size
        )

        valid = np.ones(batch_size, dtype=bool)
        for repeat in range(OUTER_REPEATS):
            for fold_id in range(1, OUTER_SPLITS + 1):
                within_fold = fold_by_repeat[repeat] == fold_id
                positive_count = weights[:, within_fold & (y == 1)].sum(axis=1)
                negative_count = weights[:, within_fold & (y == 0)].sum(axis=1)
                valid &= (positive_count > 0) & (negative_count > 0)

        if np.any(valid):
            selected = weights[valid]
            accepted.append(selected)
            n_accepted += len(selected)

    return np.vstack(accepted)[:N_BOOTSTRAP]


def weighted_auc_bootstrap(
    y_fold: np.ndarray,
    probability_fold: np.ndarray,
    weights_fold: np.ndarray,
) -> np.ndarray:
    positive = y_fold == 1
    negative = ~positive
    positive_scores = probability_fold[positive]
    negative_scores = probability_fold[negative]
    concordance = (
        (positive_scores[:, None] > negative_scores[None, :]).astype(float)
        + 0.5
        * (positive_scores[:, None] == negative_scores[None, :]).astype(float)
    )
    positive_weights = weights_fold[:, positive].astype(float, copy=False)
    negative_weights = weights_fold[:, negative].astype(float, copy=False)
    numerator = np.einsum(
        "bi,ij,bj->b",
        positive_weights,
        concordance,
        negative_weights,
        optimize=True,
    )
    denominator = positive_weights.sum(axis=1) * negative_weights.sum(axis=1)
    return numerator / denominator


def weighted_average_precision_bootstrap(
    y_fold: np.ndarray,
    probability_fold: np.ndarray,
    weights_fold: np.ndarray,
) -> np.ndarray:
    """Tie-aware weighted average precision for all bootstrap rows at once."""
    order = np.argsort(-probability_fold, kind="mergesort")
    sorted_scores = probability_fold[order]
    sorted_y = y_fold[order].astype(float)
    sorted_weights = weights_fold[:, order].astype(float, copy=False)

    group_id = np.zeros(len(order), dtype=int)
    if len(order) > 1:
        group_id[1:] = np.cumsum(sorted_scores[1:] != sorted_scores[:-1])
    n_groups = int(group_id[-1]) + 1
    group_matrix = np.zeros((len(order), n_groups), dtype=float)
    group_matrix[np.arange(len(order)), group_id] = 1.0

    group_total = sorted_weights @ group_matrix
    group_positive = (sorted_weights * sorted_y) @ group_matrix
    cumulative_total = np.cumsum(group_total, axis=1)
    cumulative_positive = np.cumsum(group_positive, axis=1)
    precision = np.divide(
        cumulative_positive,
        cumulative_total,
        out=np.zeros_like(cumulative_positive),
        where=cumulative_total > 0,
    )
    total_positive = cumulative_positive[:, -1]
    return np.sum(precision * group_positive, axis=1) / total_positive


def bootstrap_mean_metric(
    y: np.ndarray,
    prediction_by_repeat: np.ndarray,
    fold_by_repeat: np.ndarray,
    name: str,
    bootstrap_weights: np.ndarray,
) -> np.ndarray:
    """Vectorized participant-bootstrap distribution of the table statistic."""
    bootstrap_repeat_values = np.empty(
        (len(bootstrap_weights), OUTER_REPEATS), dtype=float
    )

    for repeat in range(OUTER_REPEATS):
        fold_values: list[np.ndarray] = []
        fold_weight_totals: list[np.ndarray] = []
        for fold_id in range(1, OUTER_SPLITS + 1):
            within_fold = fold_by_repeat[repeat] == fold_id
            y_fold = y[within_fold]
            probability_fold = prediction_by_repeat[repeat, within_fold]
            weights_fold = bootstrap_weights[:, within_fold]

            if name == "AUROC":
                values = weighted_auc_bootstrap(
                    y_fold, probability_fold, weights_fold
                )
            elif name == "AUPRC":
                values = weighted_average_precision_bootstrap(
                    y_fold, probability_fold, weights_fold
                )
            elif name == "Brier score":
                squared_error = (probability_fold - y_fold) ** 2
                total_weight = weights_fold.sum(axis=1).astype(float)
                values = (weights_fold @ squared_error) / total_weight
                fold_weight_totals.append(total_weight)
            else:
                raise ValueError(name)
            fold_values.append(values)

        fold_matrix = np.column_stack(fold_values)
        if name == "Brier score":
            weight_matrix = np.column_stack(fold_weight_totals)
            bootstrap_repeat_values[:, repeat] = np.average(
                fold_matrix, axis=1, weights=weight_matrix
            )
        else:
            bootstrap_repeat_values[:, repeat] = np.mean(fold_matrix, axis=1)

    return np.mean(bootstrap_repeat_values, axis=1)


def summarize_metric(
    y: np.ndarray,
    prediction_by_repeat: np.ndarray,
    fold_by_repeat: np.ndarray,
    name: str,
    bootstrap_weights: np.ndarray,
) -> str:
    repeat_values = repeat_level_fold_scores(
        y, prediction_by_repeat, fold_by_repeat, name
    )
    mean_value = float(np.mean(repeat_values))
    sd_value = float(np.std(repeat_values, ddof=1))
    bootstrap_means = bootstrap_mean_metric(
        y, prediction_by_repeat, fold_by_repeat, name, bootstrap_weights
    )
    low, high = np.quantile(bootstrap_means, [0.025, 0.975])
    return f"{mean_value:.3f} ± {sd_value:.3f} ({low:.3f}, {high:.3f})"


def make_table_s8(
    y: pd.Series,
    predictions: dict[tuple[str, str], np.ndarray],
    fold_by_repeat: np.ndarray,
) -> pd.DataFrame:
    rows: list[dict[str, str]] = []
    y_array = y.to_numpy(dtype=int)
    bootstrap_weights = make_stratified_bootstrap_weights(
        y_array, fold_by_repeat, np.random.default_rng(RANDOM_SEED)
    )
    row_order = [
        ("all_features", "catboost"),
        ("all_features", "lightgbm"),
        ("fixed_top4", "catboost"),
        ("fixed_top4", "lightgbm"),
        ("fixed_ranks5_8", "catboost"),
        ("fixed_ranks5_8", "lightgbm"),
    ]

    for panel, algo in row_order:
        matrix = predictions[(panel, algo)]
        if np.isnan(matrix).any():
            missing = int(np.isnan(matrix).sum())
            raise AssertionError(f"Missing {missing} held-out predictions for {panel}/{algo}.")
        rows.append(
            {
                "Feature set": PANEL_LABELS[panel],
                "Algorithm": ALGORITHM_LABELS[algo],
                "AUPRC mean ± SD (95% CI)": summarize_metric(
                    y_array, matrix, fold_by_repeat, "AUPRC", bootstrap_weights
                ),
                "AUROC mean ± SD (95% CI)": summarize_metric(
                    y_array, matrix, fold_by_repeat, "AUROC", bootstrap_weights
                ),
                "Brier score mean ± SD (95% CI)": summarize_metric(
                    y_array, matrix, fold_by_repeat, "Brier score", bootstrap_weights
                ),
            }
        )
    return pd.DataFrame(rows)



def style_cell(cell, *, header: bool, left: bool) -> None:
    cell.font = Font(name=ARTICLE_FONT, size=9, bold=header, color="000000")
    cell.alignment = Alignment(
        horizontal="left" if left else "center",
        vertical="center",
        wrap_text=True,
    )
    cell.border = Border(bottom=HEADER_SIDE if header else ROW_SIDE)


def write_table_s8(table: pd.DataFrame) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Table S8"
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A3"

    worksheet.append(["Feature set", "Algorithm", "AUPRC", "AUROC", "Brier score"])
    worksheet.append([None, None, "Mean ± SD (95% CI)", "Mean ± SD (95% CI)", "Mean ± SD (95% CI)"])
    worksheet.merge_cells("A1:A2")
    worksheet.merge_cells("B1:B2")

    source_columns = [
        "Feature set",
        "Algorithm",
        "AUPRC mean ± SD (95% CI)",
        "AUROC mean ± SD (95% CI)",
        "Brier score mean ± SD (95% CI)",
    ]
    for row in table[source_columns].itertuples(index=False, name=None):
        worksheet.append(list(row))

    widths = (45, 15, 24, 24, 27)
    for column_index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = width
    worksheet.row_dimensions[1].height = 28
    worksheet.row_dimensions[2].height = 30

    for row_index, row in enumerate(worksheet.iter_rows(), start=1):
        for column_index, cell in enumerate(row, start=1):
            style_cell(cell, header=row_index <= 2, left=column_index == 1 and row_index > 2)
        if row_index > 2:
            worksheet.row_dimensions[row_index].height = 23

    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_margins = PageMargins(
        left=0.25, right=0.25, top=0.35, bottom=0.35, header=0.15, footer=0.15
    )

    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    output = TABLE_DIR / "TableS8_model_performance.xlsx"
    workbook.save(output)
    print(f"[PASS] article table written: {output}")
    return output

# ---------------------------------------------------------------------------
# Main execution
# ---------------------------------------------------------------------------


def main() -> None:
    df, X, y = load_data()
    splits = load_split_manifest(df["participant_id"])
    fold_by_repeat = build_fold_assignment_matrix(df["participant_id"], splits)
    panel_features = load_fixed_feature_panels()

    LGB_MODEL_DIR.mkdir(parents=True, exist_ok=True)
    CAT_MODEL_DIR.mkdir(parents=True, exist_ok=True)
    TABLE_DIR.mkdir(parents=True, exist_ok=True)

    # Remove only the retained model sets after all required inputs validate.
    for pattern in (
        "all_features_lightgbm_repeat*_fold*.txt",
        "fixed_top4_lightgbm_repeat*_fold*.txt",
    ):
        for path in LGB_MODEL_DIR.glob(pattern):
            path.unlink()
    for pattern in (
        "all_features_catboost_repeat*_fold*.cbm",
        "fixed_top4_catboost_repeat*_fold*.cbm",
    ):
        for path in CAT_MODEL_DIR.glob(pattern):
            path.unlink()

    print("Fixed feature panels read from Table S7:")
    print("  Top-4: " + ", ".join(DISPLAY_NAME[f] for f in panel_features["fixed_top4"]))
    print(
        "  Ranks 5-8: "
        + ", ".join(DISPLAY_NAME[f] for f in panel_features["fixed_ranks5_8"])
    )
    print(
        "  All panels retain the same 122 participants and the same outer "
        "train/test assignments; tree models receive native missing values."
    )
    print(
        "  Table S8 uses fold-specific AUROC/AUPRC averaged within each repeat; "
        "Brier scores are test-size weighted within each repeat."
    )

    panel_order = ("all_features", "fixed_top4", "fixed_ranks5_8")
    combinations = [
        (panel, algo)
        for panel in panel_order
        for algo in ("lightgbm", "catboost")
    ]
    predictions = {
        combination: np.full((OUTER_REPEATS, len(df)), np.nan, dtype=float)
        for combination in combinations
    }

    for outer_fold_id in range(OUTER_SPLITS * OUTER_REPEATS):
        train_indices, test_indices, repeat_id, fold_id = outer_indices(
            df, splits, outer_fold_id
        )
        X_train, X_test = X.iloc[train_indices], X.iloc[test_indices]
        y_train = y.iloc[train_indices]

        # No endpoint-specific filtering is permitted: these row counts are
        # identical for all three panels in the current outer fold.
        expected_train_n = len(train_indices)
        expected_test_n = len(test_indices)

        for panel in panel_order:
            features = panel_features[panel]
            if len(features) != len(set(features)):
                raise AssertionError(f"Duplicate feature in panel {panel}.")
            if len(X_train) != expected_train_n or len(X_test) != expected_test_n:
                raise AssertionError("Panel-specific row filtering is not allowed.")

            for algo in ("lightgbm", "catboost"):
                probability, model, params, iterations, inner_auc = tune_and_refit(
                    algo,
                    X_train,
                    y_train,
                    X_test,
                    features,
                    outer_fold_id,
                )
                predictions[(panel, algo)][repeat_id - 1, test_indices] = probability

                if panel in {"all_features", "fixed_top4"}:
                    save_interpretation_model(panel, algo, model, repeat_id, fold_id)

                outer_auc = float(roc_auc_score(y.iloc[test_indices], probability))
                print(
                    f"[model] outer {outer_fold_id + 1:02d}/50 | "
                    f"repeat {repeat_id}, fold {fold_id} | {panel} | {algo} | "
                    f"features={len(features)} | inner AUROC={inner_auc:.3f} | "
                    f"iterations={iterations} | outer AUROC={outer_auc:.3f} | "
                    f"params={params}"
                )

    expected_models = OUTER_SPLITS * OUTER_REPEATS
    model_checks = {
        "all-feature LightGBM": len(
            list(LGB_MODEL_DIR.glob("all_features_lightgbm_repeat*_fold*.txt"))
        ),
        "fixed-Top-4 LightGBM": len(
            list(LGB_MODEL_DIR.glob("fixed_top4_lightgbm_repeat*_fold*.txt"))
        ),
        "all-feature CatBoost": len(
            list(CAT_MODEL_DIR.glob("all_features_catboost_repeat*_fold*.cbm"))
        ),
        "fixed-Top-4 CatBoost": len(
            list(CAT_MODEL_DIR.glob("fixed_top4_catboost_repeat*_fold*.cbm"))
        ),
    }
    bad = {name: count for name, count in model_checks.items() if count != expected_models}
    if bad:
        raise AssertionError(f"Expected 50 models in each retained set; observed {bad}.")

    table = make_table_s8(y, predictions, fold_by_repeat)
    output = write_table_s8(table)

    print("\n[PASS] repeated nested-CV model workflow completed")
    print(f"[PASS] LightGBM model directory: {LGB_MODEL_DIR} (100 files)")
    print(f"[PASS] CatBoost model directory: {CAT_MODEL_DIR} (100 files)")
    print(f"[PASS] Table S8: {output}")


if __name__ == "__main__":
    main()
