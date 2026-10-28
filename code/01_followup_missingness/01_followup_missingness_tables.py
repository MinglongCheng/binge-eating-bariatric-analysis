#!/usr/bin/env python3
"""Create the article-ready follow-up missingness tables without changing models.

Input:
    data/bariatric_longitudinal.csv
Outputs:
    outputs/01_followup_missingness/table/TableS1_visit_attendance.xlsx
    outputs/01_followup_missingness/table/TableS2_indicator_availability.xlsx
"""

from __future__ import annotations

import warnings
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd
import statsmodels.api as sm
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.utils import get_column_letter
from scipy.stats import chi2, fisher_exact
from statsmodels.stats.multitest import multipletests
from statsmodels.tools import add_constant
from statsmodels.tools.sm_exceptions import PerfectSeparationWarning

# Configuration
PROJECT_DIR = Path(__file__).resolve().parents[2]
DATA_FILE = PROJECT_DIR / "data" / "bariatric_longitudinal.csv"
TABLE_DIR = PROJECT_DIR / "outputs" / "01_followup_missingness" / "table"

Q_THRESHOLD = 0.10
VISITS = (1, 3, 6)
VISIT_LABELS = {1: "Month 1", 3: "Month 3", 6: "Month 6"}
COVARIATES = ("group_be", "age_z", "sex_male", "baseline_bmi_z")
COVARIATE_LABELS = {
    "group_be": "Group (BE)",
    "age_z": "Age",
    "sex_male": "Sex",
    "baseline_bmi_z": "Baseline BMI",
}

# Frozen article order; attendance-pattern flags are not clinical indicators.
INDICATORS: tuple[str, ...] = (
    "weight_kg",
    "bmi_kg_m2",
    "total_protein_g_l",
    "calcium_mmol_l",
    "vitamin_d_25oh_ng_ml",
    "iron_umol_l",
    "hemoglobin_g_l",
    "rbc_10e12_l",
    "folate_ng_ml",
    "vitamin_b12_pg_ml",
    "alt_u_l",
    "ast_u_l",
    "ggt_u_l",
    "total_cholesterol_mmol_l",
    "ldl_c_mmol_l",
    "hdl_c_mmol_l",
    "fasting_plasma_glucose_mmol_l",
    "hba1c_percent",
    "urea_mmol_l",
    "creatinine_umol_l",
    "uric_acid_umol_l",
    "sodium_mmol_l",
    "potassium_mmol_l",
    "chloride_mmol_l",
    "wbc_10e9_l",
    "neutrophils_10e9_l",
    "lymphocytes_10e9_l",
    "monocytes_10e9_l",
    "eosinophils_10e9_l",
    "basophils_10e9_l",
)

DISPLAY_NAMES = {
    "weight_kg": "Weight",
    "bmi_kg_m2": "BMI",
    "total_protein_g_l": "Total protein",
    "calcium_mmol_l": "Ca",
    "vitamin_d_25oh_ng_ml": "25(OH)D",
    "iron_umol_l": "Fe",
    "hemoglobin_g_l": "Hb",
    "rbc_10e12_l": "RBC",
    "folate_ng_ml": "Folate",
    "vitamin_b12_pg_ml": "Vitamin B12",
    "alt_u_l": "ALT",
    "ast_u_l": "AST",
    "ggt_u_l": "GGT",
    "total_cholesterol_mmol_l": "Total cholesterol",
    "ldl_c_mmol_l": "LDL-C",
    "hdl_c_mmol_l": "HDL-C",
    "fasting_plasma_glucose_mmol_l": "FPG",
    "hba1c_percent": "HbA1c",
    "urea_mmol_l": "Urea",
    "creatinine_umol_l": "Creatinine",
    "uric_acid_umol_l": "Uric acid",
    "sodium_mmol_l": "Na",
    "potassium_mmol_l": "K",
    "chloride_mmol_l": "Cl",
    "wbc_10e9_l": "WBC",
    "neutrophils_10e9_l": "Neutrophils",
    "lymphocytes_10e9_l": "Lymphocytes",
    "monocytes_10e9_l": "Monocytes",
    "eosinophils_10e9_l": "Eosinophils",
    "basophils_10e9_l": "Basophils",
}


@dataclass(frozen=True)
class ModelResult:
    lrt_p: float
    pvalues: dict[str, float]


# Data validation and model helpers

def zscore(series: pd.Series) -> pd.Series:
    values = pd.to_numeric(series, errors="coerce")
    mean = float(np.nanmean(values))
    sd = float(np.nanstd(values, ddof=0))
    if not (np.isfinite(sd) and sd > 0):
        return pd.Series(np.nan, index=series.index, dtype=float)
    return (values - mean) / sd


def normalise_group(series: pd.Series) -> pd.Series:
    values = series.astype("string").str.strip().str.lower()
    mapped = values.map({"be": 1, "non-be": 0, "non_be": 0, "nonbe": 0})
    unknown = sorted(values[mapped.isna() & values.notna()].dropna().unique().tolist())
    if unknown:
        raise ValueError(f"Unrecognised bes_group value(s): {unknown}")
    return mapped.astype(float)


def normalise_sex(series: pd.Series) -> pd.Series:
    values = series.astype("string").str.strip().str.lower()
    mapped = values.map({"male": 1, "m": 1, "female": 0, "f": 0})
    unknown = sorted(values[mapped.isna() & values.notna()].dropna().unique().tolist())
    if unknown:
        raise ValueError(f"Unrecognised sex value(s): {unknown}")
    return mapped.astype(float)


def load_data() -> pd.DataFrame:
    """Load and validate the frozen longitudinal input used by Tables S1 and S2."""
    if not DATA_FILE.exists():
        raise FileNotFoundError(f"Input file not found: {DATA_FILE}")

    df = pd.read_csv(DATA_FILE)
    df.columns = [str(column).strip() for column in df.columns]
    df = df.replace(r"^\s*$", np.nan, regex=True)

    required = {
        "participant_id",
        "sex",
        "age_years",
        "baseline_bmi_kg_m2",
        "bes_group",
        "visit_month",
        "visit_attended",
        *INDICATORS,
    }
    missing = sorted(required.difference(df.columns))
    if missing:
        raise ValueError(f"Missing required column(s): {missing}")

    numeric_columns = (
        "participant_id",
        "age_years",
        "baseline_bmi_kg_m2",
        "visit_month",
        "visit_attended",
        *INDICATORS,
    )
    for column in numeric_columns:
        df[column] = pd.to_numeric(df[column], errors="coerce")

    if df["participant_id"].isna().any() or df["visit_month"].isna().any():
        raise ValueError("participant_id and visit_month must be complete and numeric.")
    if df.duplicated(["participant_id", "visit_month"]).any():
        raise ValueError("Duplicate participant_id × visit_month rows were found.")
    if set(df["visit_month"].dropna().astype(int).unique()) != {0, 1, 3, 6}:
        raise ValueError("visit_month must contain exactly 0, 1, 3, and 6.")
    if not set(df["visit_attended"].dropna().unique()).issubset({0, 1, 0.0, 1.0}):
        raise ValueError("visit_attended must be coded 0/1.")

    df["group_be"] = normalise_group(df["bes_group"])
    df["sex_male"] = normalise_sex(df["sex"])
    df["age_z"] = zscore(df["age_years"])
    df["baseline_bmi_z"] = zscore(df["baseline_bmi_kg_m2"])
    df["visit_attended"] = df["visit_attended"].fillna(0).astype(int)

    if len(df) != 488 or df["participant_id"].nunique() != 122:
        raise AssertionError("Expected 488 rows and 122 participants.")
    baseline = df.loc[df["visit_month"].eq(0)]
    if baseline["bes_group"].value_counts().to_dict() != {"Non-BE": 66, "BE": 56}:
        raise AssertionError("Expected 66 Non-BE and 56 BE participants.")
    observed_attendance = {
        visit: int(df.loc[df["visit_month"].eq(visit), "visit_attended"].sum())
        for visit in VISITS
    }
    if observed_attendance != {1: 77, 3: 87, 6: 69}:
        raise AssertionError(f"Unexpected attendance counts: {observed_attendance}")

    print("[PASS] sample check: 488 rows; 122 participants; Non-BE=66; BE=56")
    print("[PASS] attendance check: M1=77; M3=87; M6=69")
    return df


def fit_binomial_lrt(frame: pd.DataFrame, outcome: str) -> ModelResult:
    """Fit the unchanged adjusted binomial GLM and likelihood-ratio test."""
    model_data = frame[[outcome, *COVARIATES]].dropna().copy()
    if model_data.empty or model_data[outcome].nunique() < 2:
        return ModelResult(np.nan, {})

    try:
        x = add_constant(model_data[list(COVARIATES)], has_constant="add")
        y = model_data[outcome].astype(int)
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", category=RuntimeWarning)
            warnings.filterwarnings("ignore", category=PerfectSeparationWarning)
            result = sm.GLM(y, x, family=sm.families.Binomial()).fit()
            lr_stat = float(result.null_deviance - result.deviance)
            lrt_p = float(chi2.sf(lr_stat, result.df_model))
            pvalues = {
                name: float(result.pvalues.get(name, np.nan)) for name in COVARIATES
            }
        return ModelResult(lrt_p, pvalues)
    except Exception:
        return ModelResult(np.nan, {})


def bh_fdr(values: Iterable[float]) -> np.ndarray:
    pvalues = np.asarray(list(values), dtype=float)
    safe = np.where(np.isfinite(pvalues), pvalues, 1.0)
    return multipletests(safe, method="fdr_bh")[1]


def format_count_percent(numerator: int, denominator: int) -> str:
    return f"{numerator}/{denominator} ({100.0 * numerator / denominator:.1f})"


def format_q(value: float) -> str:
    if not np.isfinite(value):
        return "NA"
    if value < 0.001:
        return "<0.001"
    return f"{value:.3f}"


def classify_adjusted_signal(
    full_q: float, covariate_q: dict[str, float]
) -> tuple[str, list[str]]:
    drivers = [
        COVARIATE_LABELS[name]
        for name in COVARIATES
        if np.isfinite(covariate_q.get(name, np.nan))
        and covariate_q[name] < Q_THRESHOLD
    ]
    if (np.isfinite(full_q) and full_q < Q_THRESHOLD) or drivers:
        return "MAR-drivers", drivers or ["Model-level signal"]
    return "MCAR-compatible", []


def fisher_group_missingness(attenders: pd.DataFrame, variable: str) -> float:
    """Retain the original Fisher fallback for Table S2 group missingness."""
    temp = attenders[["group_be", variable]].copy()
    temp["missing"] = temp[variable].isna().astype(int)
    table = pd.crosstab(temp["group_be"], temp["missing"])
    if not ({0.0, 1.0} <= set(table.index)) or not ({0, 1} <= set(table.columns)):
        return np.nan
    matrix = [
        [int(table.loc[0.0, 0]), int(table.loc[0.0, 1])],
        [int(table.loc[1.0, 0]), int(table.loc[1.0, 1])],
    ]
    return float(fisher_exact(matrix, alternative="two-sided")[1])


# Table calculations

def build_table_s1(df: pd.DataFrame) -> pd.DataFrame:
    models: dict[int, ModelResult] = {}
    for visit in VISITS:
        subset = df.loc[df["visit_month"].eq(visit)].copy()
        subset["no_show"] = 1 - subset["visit_attended"]
        models[visit] = fit_binomial_lrt(subset, "no_show")

    full_q = bh_fdr(models[visit].lrt_p for visit in VISITS)
    covariate_q = {
        covariate: bh_fdr(
            models[visit].pvalues.get(covariate, np.nan) for visit in VISITS
        )
        for covariate in COVARIATES
    }

    rows: list[dict[str, object]] = []
    for index, visit in enumerate(VISITS):
        subset = df.loc[df["visit_month"].eq(visit)]
        be = subset["group_be"].eq(1)
        non_be = subset["group_be"].eq(0)

        n_all = len(subset)
        n_be = int(be.sum())
        n_non_be = int(non_be.sum())
        attended_all = int(subset["visit_attended"].sum())
        attended_be = int(subset.loc[be, "visit_attended"].sum())
        attended_non_be = int(subset.loc[non_be, "visit_attended"].sum())

        group_gap = abs(
            100.0 * attended_be / n_be - 100.0 * attended_non_be / n_non_be
        )
        q_full = float(full_q[index])
        q_covariates = {
            covariate: float(covariate_q[covariate][index])
            for covariate in COVARIATES
        }
        classification, drivers = classify_adjusted_signal(q_full, q_covariates)
        minimum_q = float(np.nanmin([q_full, *q_covariates.values()]))

        rows.append(
            {
                "Visit": VISIT_LABELS[visit],
                "All n/N (%)": format_count_percent(attended_all, n_all),
                "BE n/N (%)": format_count_percent(attended_be, n_be),
                "Non-BE n/N (%)": format_count_percent(
                    attended_non_be, n_non_be
                ),
                "Absolute group difference (pp)": f"{group_gap:.1f}",
                "Classification": classification,
                "Drivers": ", ".join(drivers) if drivers else "-",
                "Minimum q (BH-FDR)": format_q(minimum_q),
            }
        )
    return pd.DataFrame(rows)


def build_table_s2(df: pd.DataFrame) -> pd.DataFrame:
    records: list[dict[str, object]] = []
    for visit in VISITS:
        attenders = df.loc[
            df["visit_month"].eq(visit) & df["visit_attended"].eq(1)
        ].copy()
        be = attenders["group_be"].eq(1)
        non_be = attenders["group_be"].eq(0)

        for variable in INDICATORS:
            model_frame = attenders[[variable, *COVARIATES]].copy()
            model_frame["indicator_missing"] = model_frame[variable].isna().astype(int)
            model = fit_binomial_lrt(model_frame, "indicator_missing")
            records.append(
                {
                    "visit_month": visit,
                    "variable": variable,
                    "n_all": len(attenders),
                    "available_all": int(attenders[variable].notna().sum()),
                    "n_be": int(be.sum()),
                    "available_be": int(attenders.loc[be, variable].notna().sum()),
                    "n_non_be": int(non_be.sum()),
                    "available_non_be": int(
                        attenders.loc[non_be, variable].notna().sum()
                    ),
                    "full_model_p": model.lrt_p,
                    "fisher_group_p": fisher_group_missingness(attenders, variable),
                    **{
                        f"{covariate}_p": model.pvalues.get(covariate, np.nan)
                        for covariate in COVARIATES
                    },
                }
            )

    details = pd.DataFrame(records)
    for visit in VISITS:
        mask = details["visit_month"].eq(visit)
        details.loc[mask, "full_model_q"] = bh_fdr(
            details.loc[mask, "full_model_p"]
        )
        for covariate in COVARIATES:
            details.loc[mask, f"{covariate}_q"] = bh_fdr(
                details.loc[mask, f"{covariate}_p"]
            )

    visit_types: list[str] = []
    for _, row in details.iterrows():
        q_covariates = {
            covariate: float(row[f"{covariate}_q"]) for covariate in COVARIATES
        }
        adjusted_type, _ = classify_adjusted_signal(
            float(row["full_model_q"]), q_covariates
        )
        if adjusted_type == "MAR-drivers":
            visit_type = "MAR-drivers"
        elif np.isfinite(row["fisher_group_p"]) and float(row["fisher_group_p"]) < 0.05:
            visit_type = "MAR-G"
        else:
            visit_type = "MCAR-compatible"
        visit_types.append(visit_type)
    details["visit_type"] = visit_types

    rows: list[dict[str, object]] = []
    for variable in INDICATORS:
        variable_rows = details.loc[details["variable"].eq(variable)].set_index(
            "visit_month"
        )
        visit_cells: dict[int, str] = {}
        gaps: list[float] = []
        types: list[str] = []

        for visit in VISITS:
            row = variable_rows.loc[visit]
            visit_cells[visit] = " | ".join(
                (
                    f"{int(row['available_all'])}/{int(row['n_all'])}",
                    f"{int(row['available_be'])}/{int(row['n_be'])}",
                    f"{int(row['available_non_be'])}/{int(row['n_non_be'])}",
                )
            )
            be_percent = 100.0 * float(row["available_be"]) / float(row["n_be"])
            non_be_percent = (
                100.0 * float(row["available_non_be"]) / float(row["n_non_be"])
            )
            gaps.append(abs(be_percent - non_be_percent))
            types.append(str(row["visit_type"]))

        if "MAR-drivers" in types:
            overall_type = "MAR-drivers"
        elif "MAR-G" in types:
            overall_type = "MAR-G"
        else:
            overall_type = "MCAR-compatible"

        rows.append(
            {
                "Indicator": DISPLAY_NAMES[variable],
                "M1 availability": visit_cells[1],
                "M3 availability": visit_cells[3],
                "M6 availability": visit_cells[6],
                "Maximum group difference (pp)": f"{max(gaps):.1f}",
                "Any MAR signal": "Yes"
                if overall_type != "MCAR-compatible"
                else "No",
                "Overall classification": overall_type,
            }
        )

    return pd.DataFrame(rows)


# XLSX export

THIN_GRAY = Side(style="thin", color="B7B7B7")
THIN_BLACK = Side(style="thin", color="000000")
BODY_BORDER = Border(bottom=THIN_GRAY)
HEADER_BORDER = Border(top=THIN_BLACK, bottom=THIN_BLACK)
FONT_BODY = Font(name="Arial", size=9)
FONT_HEADER = Font(name="Arial", size=9, bold=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_data_sheet(
    worksheet,
    *,
    header_rows: int,
    column_widths: tuple[float, ...],
) -> None:
    """Apply the frozen article-table formatting without decorative colour fills."""
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = f"A{header_rows + 1}"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.orientation = worksheet.ORIENTATION_LANDSCAPE
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.page_margins = PageMargins(
        left=0.25, right=0.25, top=0.40, bottom=0.40, header=0.15, footer=0.15
    )
    worksheet.print_options.horizontalCentered = True
    worksheet.print_area = f"A1:{get_column_letter(len(column_widths))}{worksheet.max_row}"
    worksheet.print_title_rows = f"1:{header_rows}"

    for column_index, width in enumerate(column_widths, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    for row in worksheet.iter_rows():
        for cell in row:
            is_header = cell.row <= header_rows
            cell.font = FONT_HEADER if is_header else FONT_BODY
            cell.alignment = ALIGN_LEFT if cell.column == 1 else ALIGN_CENTER
            cell.border = HEADER_BORDER if is_header else BODY_BORDER

    for row_index in range(1, header_rows + 1):
        worksheet.row_dimensions[row_index].height = 32
    for row_index in range(header_rows + 1, worksheet.max_row + 1):
        worksheet.row_dimensions[row_index].height = 18


def write_table_s1(table: pd.DataFrame, output_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Table S1"

    worksheet.append(table.columns.tolist())
    for row in table.itertuples(index=False, name=None):
        worksheet.append(list(row))

    style_data_sheet(
        worksheet,
        header_rows=1,
        column_widths=(12, 17, 17, 19, 23, 19, 29, 20),
    )
    workbook.save(output_path)
    print(f"[PASS] article table written: {output_path}")


def write_table_s2(table: pd.DataFrame, output_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Table S2"

    top_header = [
        "Indicator",
        "M1 availability, n/N",
        "M3 availability, n/N",
        "M6 availability, n/N",
        "Maximum group difference (pp)",
        "Any MAR signal",
        "Overall classification",
    ]
    sub_header = [
        None,
        "All | BE | Non-BE",
        "All | BE | Non-BE",
        "All | BE | Non-BE",
        None,
        None,
        None,
    ]
    worksheet.append(top_header)
    worksheet.append(sub_header)
    for row in table.itertuples(index=False, name=None):
        worksheet.append(list(row))

    for cell_range in ("A1:A2", "E1:E2", "F1:F2", "G1:G2"):
        worksheet.merge_cells(cell_range)

    style_data_sheet(
        worksheet,
        header_rows=2,
        column_widths=(24, 25, 25, 25, 23, 17, 22),
    )
    workbook.save(output_path)
    print(f"[PASS] article table written: {output_path}")


def main() -> None:
    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    df = load_data()
    table_s1 = build_table_s1(df)
    table_s2 = build_table_s2(df)

    if table_s1.shape != (3, 8):
        raise AssertionError(f"Unexpected Table S1 shape: {table_s1.shape}")
    if table_s2.shape != (30, 7):
        raise AssertionError(f"Unexpected Table S2 shape: {table_s2.shape}")

    write_table_s1(table_s1, TABLE_DIR / "TableS1_visit_attendance.xlsx")
    write_table_s2(table_s2, TABLE_DIR / "TableS2_indicator_availability.xlsx")
    print("[PASS] module 01 table workflow completed")


if __name__ == "__main__":
    main()
